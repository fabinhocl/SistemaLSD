from pyexpat.errors import messages
from dal import autocomplete
from AppLSD.models import Family, Aluno, Turma, Activity, FrequenciaTurma, FrequenciaAluno, FrequenciaAtividade, MovimentacaoTurmaAluno, OcorrenciaAluno, Adult, AppLog, PerfilUsuario, DocumentoFamilia    
from AppLSD.utils import is_coordenacao, is_educadora, coordenacao_required, usuario_tem_perfil, get_tipo_perfil, registrar_log, TIPOS_PERFIL_VALIDOS, sincronizar_grupos_usuario, pode_editar_frequencia_turma # ✅ Importar as funções
from AppLSD.templatetags.perfil_tags import has_perfil
from .forms import FamilyForm, AlunoForm, TurmaForm, ActivityForm, AddAlunosToTurmaForm, AlunoFiltroForm, AlunoInlineFormSet, MoverAlunoForm, OcorrenciaAlunoForm, AdultFormSet, AdultForm, UsuarioCadastroForm, UsuarioEdicaoForm, RemoverAlunoAtividadeForm, MeuPerfilForm
from calendar import monthrange
from collections import defaultdict, OrderedDict
from django import forms
from django.core.exceptions import PermissionDenied
from django.core.files.storage import default_storage
from django.core.paginator import Paginator
from django.contrib import messages
from django.contrib.admin.views.decorators import staff_member_required
from django.contrib.auth import authenticate 
from django.contrib.auth.mixins import LoginRequiredMixin
from django.contrib.auth.decorators import login_required, user_passes_test
from django.contrib.auth.hashers import check_password
from django.contrib.auth.models import User, Group
from django.contrib.contenttypes.models import ContentType
from django.contrib.staticfiles import finders
from django.shortcuts import render, redirect, get_object_or_404
from django.utils import timezone
from django.utils.timezone import make_aware
from django.http import JsonResponse, HttpResponse, HttpResponseForbidden, Http404
from django.db import transaction
from django.db.models import Count, Q, F, Avg, Value, When, Sum, CharField, Exists, OuterRef, IntegerField, FloatField, Case
from django.db.models.functions import Cast, Coalesce, Concat, Lower
from django.forms.models import inlineformset_factory
from django.template.loader import render_to_string
from django.views.generic import ListView
from django_xhtml2pdf.utils import generate_pdf
from datetime import date, timedelta, datetime, time
from .permissoes import require_perfil
# Concatena e ordena por data
from itertools import chain
from io import BytesIO
from operator import itemgetter
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os
os.environ['WEASYPRINT_DLL_DIRECTORIES'] = r"C:\Program Files\GTK3-Runtime Win64\bin"
from weasyprint import HTML, CSS
import plotly.graph_objs as go
import plotly.offline as opy
import io
import base64
import matplotlib.pyplot as plt
import unicodedata
import logging
import re

logger = logging.getLogger(__name__)

def root_redirect(request):
    if request.user.is_authenticated:
        return redirect('/home/')  # URL direta, não por nome
    else:
        return redirect('/accounts/login/')  # nome padrão da view de login
"""   
def require_perfil(perfil_tipo):
    def decorator(view_func):
        @wraps(view_func)
        def _wrapped_view(request, *args, **kwargs):
            if not request.user.is_authenticated:
                return HttpResponseForbidden("Não autenticado.")
            if not request.user.perfis.filter(tipo=perfil_tipo).exists():
                return HttpResponseForbidden("Você não possui permissão para esta página.")
            return view_func(request, *args, **kwargs)
        return _wrapped_view
    return decorator
"""   
@require_perfil('educadora')
def view_educadora(request):
    # apenas educadora acessa
    pass

@require_perfil('facilitador')
def view_facilitador(request):
    # apenas facilitador acessa
    pass
@require_perfil('servicosocial')
def view_servicosocial(request):
    # apenas serviço social acessa
    pass

@require_perfil('admin')
def view_admin(request):
    # apenas admin acessa
    pass

@require_perfil('educadora')
def registrar_frequencia(request, turma_id):
    perfil = request.user.perfis.get(tipo_perfil='educadora')
    turma = get_object_or_404(Turma, id=turma_id, educadora=perfil)
    # Resto da lógica

@require_perfil('servicosocial')
def registrar_frequencia(request):
    perfil = request.user.perfis.get(tipo_perfil='servicosocial')
    pass
    # Resto da lógica

@require_perfil('facilitador')
def registrar_frequencia(request, turma_id):
    perfil = request.user.perfis.get(tipo_perfil='facilitador')
    turma = get_object_or_404(Turma, id=turma_id, facilitador=perfil)
    # Resto da lógica

@require_perfil('coordenacao')
def dashboard_coordenador(request):
    # Coordenação pode acessar dashboards/admin de turmas
    pass


@require_perfil('diretoria')
def dashboard_diretoria(request):
    # Diretoria pode acessar dashboards/admin de turmas
    pass

@require_perfil('admin')
def usuarios_gerenciar(request):
    # Lista todos os usuários do sistema, com seus perfis associados
    termo = request.GET.get("busca", "")
    usuarios = User.objects.all().prefetch_related('perfis').order_by('username')
    if termo:
        usuarios = usuarios.filter(username__icontains=termo) | usuarios.filter(first_name__icontains=termo)

    # processamento POST igual ao seu código
    if request.method == "POST":
        user_id = request.POST.get("desativar_id")
        if user_id:
            usuario = get_object_or_404(User, id=user_id)
            usuario.is_active = False
            usuario.save()
            return redirect('usuarios_gerenciar')
        user_id = request.POST.get("ativar_id")
        if user_id:
            usuario = get_object_or_404(User, id=user_id)
            usuario.is_active = True
            usuario.save()
            return redirect('usuarios_gerenciar')

    # Lista de tipos de perfil a partir do model
    #from AppLSD.models import PerfilUsuario
    tipos_perfil = [tp[0] for tp in PerfilUsuario._meta.get_field('tipo_perfil').choices]

    # Se estiver mostrando perfis de um único usuário (por exemplo, ao editar), passe assim:
    perfis_do_usuario = []
    usuario_id = request.GET.get("usuario_id")  # exemplo para edição de um usuário
    if usuario_id:
        usuario = get_object_or_404(User, id=usuario_id)
        perfis_do_usuario = [perfil.tipo_perfil for perfil in usuario.perfis.all()]

    return render(request, 'AppLSD/usuarios_gerenciar.html', {
        "usuarios": usuarios,
        "tipos_perfil": tipos_perfil,
        "perfis_do_usuario": perfis_do_usuario,
    })
    

@require_perfil('admin')
def editar_perfis_usuario(request, usuario_id):
    usuario = get_object_or_404(User, id=usuario_id)

    if request.method == "POST":
        novos_perfis = request.POST.getlist('tipo_perfil')

        # mantém só tipos válidos
        novos_perfis = [tp for tp in novos_perfis if tp in TIPOS_PERFIL_VALIDOS]

        # remove perfis não marcados
        usuario.perfis.exclude(tipo_perfil__in=novos_perfis).delete()

        # adiciona os novos perfis
        for tipo in novos_perfis:
            if not usuario.perfis.filter(tipo_perfil=tipo).exists():
                PerfilUsuario.objects.create(user=usuario, tipo_perfil=tipo)

        # se ficou sem nenhum, força colaborador
        if not usuario.perfis.exists():
            PerfilUsuario.objects.create(user=usuario, tipo_perfil='colaborador')

        sincronizar_grupos_usuario(usuario)
        return redirect('usuarios_gerenciar')

    perfis_do_usuario = list(usuario.perfis.values_list('tipo_perfil', flat=True))
    return render(
        request,
        'AppLSD/editar_perfis.html',
        {
            "usuario": usuario,
            "tipo_perfil": TIPOS_PERFIL_VALIDOS,
            "perfis_do_usuario": perfis_do_usuario,
        },
    )

@require_perfil('admin')
def editar_usuario(request, usuario_id):
    usuario = get_object_or_404(User, id=usuario_id)
    #tipo_perfil_choices = TIPOS_PERFIL_VALIDOS
    perfis_do_usuario = list(usuario.perfis.values_list('tipo_perfil', flat=True))

    if request.method == "POST":
        form = UsuarioEdicaoForm(request.POST, instance=usuario)
        if form.is_valid():
            usuario = form.save()

            novos_perfis = request.POST.getlist('tipo_perfil')
            novos_perfis = [tp for tp in novos_perfis if tp in TIPOS_PERFIL_VALIDOS]

            usuario.perfis.exclude(tipo_perfil__in=novos_perfis).delete()

            for tipo in novos_perfis:
                if not usuario.perfis.filter(tipo_perfil=tipo).exists():
                    PerfilUsuario.objects.create(user=usuario, tipo_perfil=tipo)

            if not usuario.perfis.exists():
                PerfilUsuario.objects.create(user=usuario, tipo_perfil='colaborador')

            sincronizar_grupos_usuario(usuario)
            return redirect('usuarios_gerenciar')
    else:
        form = UsuarioEdicaoForm(instance=usuario)

    return render(request, 'AppLSD/editar_usuario.html', {
        'form': form,
        'usuario': usuario,
        'tipo_perfil_choices': TIPOS_PERFIL_VALIDOS,
        'perfis_do_usuario': perfis_do_usuario,
    })

@login_required
def editar_meu_perfil(request):
    if request.method == 'POST':
        form = MeuPerfilForm(request.POST, instance=request.user)
        if form.is_valid():
            form.save()
            return redirect('home')
    else:
        form = MeuPerfilForm(instance=request.user)

    return render(request, 'AppLSD/editar_meu_perfil.html', {'form': form})



@require_perfil('admin')
def resetar_senha_usuario(request, usuario_id):
    usuario = User.objects.get(id=usuario_id)
    if request.method == "POST":
        nova_senha = request.POST.get("nova_senha")
        usuario.set_password(nova_senha)
        usuario.save()
        return redirect('usuarios_gerenciar')
    return render(request, 'AppLSD/resetar_senha.html', {"usuario": usuario})

@require_perfil('admin')
def cadastrar_usuario(request):
    if request.method == 'POST':
        form = UsuarioCadastroForm(request.POST)
        if form.is_valid():
            user = form.save()
            perfis = request.POST.getlist("perfis")  # nomes do tipo_perfil
            perfis = [p for p in perfis if p in TIPOS_PERFIL_VALIDOS]
            
            if perfis:
                for perfil in perfis:
                    PerfilUsuario.objects.create(user=user, tipo_perfil=perfil)
            else:
                PerfilUsuario.objects.create(user=user, tipo_perfil='colaborador')
            sincronizar_grupos_usuario(user)
            return redirect('usuarios_gerenciar')
    else:
        form = UsuarioCadastroForm()
    return render(
        request,
        'AppLSD/cadastrar_usuario.html',
        {"form": form, "tipo_perfil": TIPOS_PERFIL_VALIDOS},
    )



@login_required
def home(request):
    if usuario_tem_perfil(request.user, "facilitador"):
        return redirect('home_facilitador')
    elif usuario_tem_perfil(request.user, "educadora"):
        return redirect('home_educadora')
    elif usuario_tem_perfil(request.user, "coordenacao"):
        return redirect('home_escola')
    elif usuario_tem_perfil(request.user, "servicosocial"):
        return redirect('home_assist')
    elif usuario_tem_perfil(request.user, "diretoria"):
        return redirect('home_diretoria')  # ou o nome correto do path para dashboard
    return render(request, "AppLSD/home.html", {
        "usuario_admin": usuario_tem_perfil(request.user, "admin"),
        "usuario_coord": usuario_tem_perfil(request.user, "coordenacao"),
        "usuario_supervisor": usuario_tem_perfil(request.user, "supervisor"),
        "usuario_educadora": usuario_tem_perfil(request.user, "educadora"),
        "usuario_servicosocial": usuario_tem_perfil(request.user, "servicosocial"),
        "usuario_facilitador": usuario_tem_perfil(request.user, "facilitador"),
        "usuario_dir": usuario_tem_perfil(request.user, "diretoria"),
        "usuario_adm": usuario_tem_perfil(request.user, "administrativo"),
        "usuario_colab": usuario_tem_perfil(request.user, "colaborador"),
        "usuario_finan": usuario_tem_perfil(request.user, "financeiro"),
        "usuario_nutri": usuario_tem_perfil(request.user, "nutricao"),

    })
   
def remove_accents(text):
    return ''.join(
        c for c in unicodedata.normalize('NFD', text)
        if unicodedata.category(c) != 'Mn'
    )


from django.contrib.auth.decorators import login_required
from django.db.models import Exists, OuterRef
from django.shortcuts import render
from django.utils import timezone

@login_required
def home_educadora(request):
    educadora = request.user
    hoje = timezone.now().date()
    weekday = hoje.weekday()

    mapa_weekday = {
        0: 'segunda',
        1: 'terca',
        2: 'quarta',
        3: 'quinta',
        4: 'sexta',
    }
    valor_dia = mapa_weekday.get(weekday)

    turmas = Turma.objects.filter(educadora=educadora).order_by('turno', 'grupo')

    turmas_com_status = []
    for turma in turmas:
        frequencia_turma_hoje = FrequenciaTurma.objects.filter(
            turma=turma,
            data=hoje
        ).first()

        total_alunos = Aluno.objects.filter(turma=turma).count()

        freq_alunos = FrequenciaAluno.objects.filter(
            chamada__turma=turma,
            chamada__data=hoje
        )

        alunos_presentes_hoje = freq_alunos.filter(presente=True).count()

        item_turma = {
            'turma': turma,
            'total_alunos': total_alunos,
            'alunos_presentes': alunos_presentes_hoje,
            'frequencia_existe': frequencia_turma_hoje is not None,
            'frequencia': frequencia_turma_hoje,
            'pode_editar': is_coordenacao(educadora),
            'pode_visualizar': is_educadora(educadora) and frequencia_turma_hoje is not None,
        }
        turmas_com_status.append(item_turma)

    alunos_da_educadora = Aluno.objects.filter(
        turma__educadora=educadora
    ).distinct()

    atividades_com_alunos_da_educadora = Activity.objects.filter(
        Exists(
            alunos_da_educadora.filter(atividades=OuterRef('pk'))
        )
    ).order_by('turno', 'atividade')

    if valor_dia:
        atividades_com_alunos_da_educadora = atividades_com_alunos_da_educadora.filter(
            dia_semana__icontains=valor_dia
        )

    atividades_com_status = []
    for atividade in atividades_com_alunos_da_educadora:
        frequencia_existe = FrequenciaAtividade.objects.filter(
            atividade=atividade,
            data=hoje,
            aluno__in=alunos_da_educadora
        ).exists()

        atividades_com_status.append({
            'obj': atividade,
            'frequencia_existe': frequencia_existe
        })

    turmas_manha = [
        item for item in turmas_com_status
        if item['turma'].turno in ['M', 'Matutino', 'matutino']
    ]

    turmas_tarde = [
        item for item in turmas_com_status
        if item['turma'].turno in ['V', 'Vespertino', 'vespertino']
    ]

    atividades_manha = [
        item for item in atividades_com_status
        if item['obj'].turno in ['M', 'Matutino', 'matutino']
    ]

    atividades_tarde = [
        item for item in atividades_com_status
        if item['obj'].turno in ['V', 'Vespertino', 'vespertino']
    ]

    context = {
        'turmas_manha': turmas_manha,
        'turmas_tarde': turmas_tarde,
        'atividades_manha': atividades_manha,
        'atividades_tarde': atividades_tarde,
        'is_coordenacao': is_coordenacao(educadora),
        'is_educadora': is_educadora(educadora),
    }

    return render(request, 'AppLSD/home_educadora.html', context)

    
@login_required
def home_facilitador(request):
    facilitador_atividades = Activity.objects.filter(facilitador=request.user)
    return render(request, 'AppLSD/home_facilitador.html', {'atividades': facilitador_atividades})

@login_required
def home_assist(request):
    #print(">>> ENTROU NA HOME_ASSIST")
    servicosocial = request.user
    hoje = timezone.now().date()
    data_corte = date(hoje.year - 60, hoje.month, hoje.day)
    contexto = {
        "total_assistidos": Aluno.objects.filter(status_lsd="Frequentando").count(),
        "total_familias": Family.objects.count(),
        "total_familias_ativas": Family.objects.filter(status="ativo").count(),
        "total_adultos": Adult.objects.count(),
        "total_idosos": Adult.objects.filter(birth_date__lte=data_corte).count(),
        # "percentual_presenca": calcular_presenca_hoje(),
        "hoje": timezone.now().strftime("%d/%m/%Y"),
    }
    #hoje = timezone.now().date()
    return render(request, 'AppLSD/home_assist.html', contexto)

@login_required
def home_escola(request):

    hoje = date.today()

    # todos aniversariantes do mês (todas as turmas)
    aniversariantes_mes = (
        Aluno.objects.select_related('turma').filter(birth_date__month=hoje.month, status_lsd="Frequentando")
        .order_by('birth_date__day', 'name')
    )

    # aniversariantes do dia (subset do anterior)
    aniversariantes_dia = aniversariantes_mes.filter(birth_date__day=hoje.day)

    context = {
        'aniversariantes_mes': aniversariantes_mes,
        'aniversariantes_dia': aniversariantes_dia,
    }
    return render(request, 'AppLSD/home_escola.html', context)

@login_required
def home_diretoria(request):
    diretoria = request.user
    hoje = timezone.now().date()
    data_corte = date(hoje.year - 60, hoje.month, hoje.day)
    contexto = {
        "total_assistidos": Aluno.objects.filter(status_lsd="Frequentando").count(),
        "total_familias": Family.objects.count(),
        "total_familias_ativas": Family.objects.filter(status="ativo").count(),
        "total_adultos": Adult.objects.count(),
        "total_idosos": Adult.objects.filter(birth_date__lte=data_corte).count(),
        # "percentual_presenca": calcular_presenca_hoje(),
        "hoje": timezone.now().strftime("%d/%m/%Y"),
    }
    #hoje = timezone.now().date()
    return render(request, 'AppLSD/home_diretoria.html', contexto)
       
@login_required
def home_adm(request):
    return render(request, 'AppLSD/home_adm.html')

#Mesclagem das views de dashboard de presença e relatórios
@login_required
def dashboard_completo(request):
    today = timezone.now().date()

    turmas = Turma.objects.all()
    atividades = Activity.objects.all()
    alunos = Aluno.objects.all()

    total_alunos = Aluno.objects.filter(status_lsd='Frequentando').count()
    
    alunos_busca = Aluno.objects.filter(status_lsd='Frequentando').order_by('name')

    presentes_hoje_manha = FrequenciaAluno.objects.filter(
        aluno__turma__turno='Matutino',
        chamada__data=today,
        presente=True
    ).count()

    presentes_hoje_tarde = FrequenciaAluno.objects.filter(
        aluno__turma__turno='Vespertino',
        chamada__data=today,
        presente=True
    ).count()

    faltas_hoje = FrequenciaAluno.objects.filter(
        chamada__data=today,
        presente=False
    ).count()

    turmas_ativas = Turma.objects.count()
    atividades_ativas = Activity.objects.count()

    # médias últimos 7 dias (se ainda quiser usar nos atalhos)
    turmas_com_frequencia = []
    for turma in Turma.objects.annotate(total_alunos=Count("alunos")):
        if turma.total_alunos > 0:
            media_presenca = (
                FrequenciaAluno.objects
                .filter(
                    aluno__turma=turma,
                    chamada__data__gte=today - timezone.timedelta(days=7),
                )
                .annotate(presente_int=Cast("presente", IntegerField()))
                .aggregate(media=Avg("presente_int"))["media"] or 0
            )
            turma.media_presenca = media_presenca * 100
            turmas_com_frequencia.append(turma)

    atividades_com_frequencia = []
    for atividade in Activity.objects.annotate(total_alunos=Count("alunos")):
        if atividade.total_alunos > 0:
            media_presenca = (
                FrequenciaAluno.objects
                .filter(
                    aluno__atividades=atividade,
                    chamada__data__gte=today - timezone.timedelta(days=7),
                )
                .annotate(presente_int=Cast("presente", IntegerField()))
                .aggregate(media=Avg("presente_int"))["media"] or 0
            )
            atividade.media_presenca = media_presenca * 100
            atividades_com_frequencia.append(atividade)

    context = {
        "total_alunos": total_alunos,
        "presentes_hoje_manha": presentes_hoje_manha,
        "presentes_hoje_tarde": presentes_hoje_tarde,
        "faltas_hoje": faltas_hoje,
        "turmas_ativas": turmas_ativas,
        "atividades_ativas": atividades_ativas,
        "turmas_com_frequencia": turmas_com_frequencia,
        "atividades_com_frequencia": atividades_com_frequencia,
        "turmas": turmas,
        "atividades": atividades,
        "alunos": alunos,
        "alunos_busca": alunos_busca,
    }
    return render(request, "AppLSD/home_relatorios.html", context)

    
#Tela para Educadora e Facilitador
@login_required
def dashboard_presenca(request):
    hoje = timezone.now().date()
    

    total_alunos = Aluno.objects.filter(status_lsd='Frequentando').count()

    turmas_ativas = Turma.objects.filter(alunos__isnull=False).distinct().count()

    presentes_hoje = FrequenciaAluno.objects.filter(
        chamada__data=hoje, presente=True
    ).count()
    
    faltas_hoje = FrequenciaAluno.objects.filter(
        chamada__data=hoje, presente=False
    ).count()

    presentes_manha = FrequenciaAluno.objects.filter(
        chamada__data=hoje,
        chamada__turma__turno__iexact='Matutino',
        presente=True
    ).count()

    faltas_manha = FrequenciaAluno.objects.filter(
        chamada__data=hoje,
        chamada__turma__turno__iexact='Matutino',
        presente=False
    ).count()

    presentes_tarde = FrequenciaAluno.objects.filter(
        chamada__data=hoje,
        chamada__turma__turno__iexact='Vespertino',
        presente=True
    ).count()

    faltas_tarde = FrequenciaAluno.objects.filter(
        chamada__data=hoje,
        chamada__turma__turno__iexact='Vespertino',
        presente=False
    ).count()
    

    # Frequência por turma
    turmas_com_frequencia = []
    for turma in (
        Turma.objects
        .annotate(
            total_alunos=Count(
                'alunos',
                filter=Q(alunos__status_lsd='Frequentando'),
                distinct=True
            )
        )
        .filter(total_alunos__gt=0)
        .order_by('turno', 'grupo')
    ):
        turma.presentes_hoje = FrequenciaAluno.objects.filter(
            chamada__turma=turma,
            chamada__data=hoje,
            presente=True
        ).count()

        turma.faltas_hoje = FrequenciaAluno.objects.filter(
            chamada__turma=turma,
            chamada__data=hoje,
            presente=False
        ).count()

        turmas_com_frequencia.append(turma)

    atividades_com_frequencia = []
    for atividade in (
        Activity.objects
        .annotate(
            total_alunos=Count('alunos', distinct=True)
        )
        .filter(total_alunos__gt=0)
        .order_by('turno', 'atividade')
    ):
        atividade.presentes_hoje = FrequenciaAtividade.objects.filter(
            atividade=atividade,
            data=hoje,
            presente=True
        ).count()

        atividade.faltas_hoje = FrequenciaAtividade.objects.filter(
            atividade=atividade,
            data=hoje,
            presente=False
        ).count()

        atividades_com_frequencia.append(atividade)

    return render(
        request,
        "AppLSD/dashboard_presenca.html",
        {
            "total_alunos": total_alunos,
            "turmas_ativas": turmas_ativas,
            "presentes_hoje": presentes_hoje,
            "faltas_hoje": faltas_hoje,
            "presentes_manha": presentes_manha,
            "faltas_manha": faltas_manha,
            "presentes_tarde": presentes_tarde,
            "faltas_tarde": faltas_tarde,
            "turmas_com_frequencia": turmas_com_frequencia,
            "atividades_com_frequencia": atividades_com_frequencia,
        },
    )

#Métódos Familia
@login_required
def family_detail(request, pk):
    family = get_object_or_404(Family, pk=pk)

    from django.contrib.contenttypes.models import ContentType
    ct = ContentType.objects.get_for_model(Family)
    logs = AppLog.objects.filter(content_type=ct, object_id=family.pk)

    ano_atual = timezone.now().year
    doc_atual = DocumentoFamilia.objects.filter(family=family, ano=ano_atual).first()

    # todos os documentos da família, mais recentes primeiro
    documentos = family.documentos.all().order_by('-criado_em')

    context = {
        "family": family,
        "logs": logs,
        "doc_atual": doc_atual,
        "documentos": documentos,
    }

    return render(request, 'AppLSD/family_detail.html', context)


def upload_documento(request, pk):
    family = get_object_or_404(Family, pk=pk)

    if request.method == "POST":
        form = DocumentoFamiliaForm(request.POST, request.FILES)
        if form.is_valid():
            doc = form.save(commit=False)
            doc.family = family
            doc.ano = timezone.now().year
            doc.save()
            return redirect("family_detail", pk=family.pk)
    else:
        form = DocumentoFamiliaForm()

    return render(request, "AppLSD/upload_documento.html", {"family": family, "form": form})


@login_required
def family_delete_confirm(request, pk):
    family = get_object_or_404(Family, pk=pk)
    error = None

    if request.method == 'POST':
        senha = request.POST.get('senha')
        usuario = request.user
        user_autenticado = authenticate(username=usuario.username, password=senha)
        if user_autenticado is not None:
            family.delete()
            return redirect('family_list')
        else:
            error = "Senha incorreta."

    return render(request, 'AppLSD/family_confirm_delete.html', {
        'family': family,
        'error': error,
    })

@login_required
def family_list(request):
    # termo vindo da URL ?q=...
    query = request.GET.get('q', '').strip()
    show_ativas = request.GET.get('show_ativas')  # 'on' se checkbox marcado
    # queryset base
    families_qs = Family.objects.all()

    # aplica filtro se houver busca
    if query:
        families_qs = families_qs.filter(
            Q(registration_number__icontains=query) |
            Q(responsible_name__icontains=query) |
            Q(cpf__icontains=query)
        )

    # filtro de status ativo
    if show_ativas == 'on':
        families_qs = families_qs.filter(status='ativo')

    # ordenação
    families_qs = families_qs.order_by('registration_number')

    # guarda total ANTES da paginação
    total = families_qs.count()
    total_ativas = families_qs.filter(status='ativo').count()

    # paginação
    paginator = Paginator(families_qs, 50)  # 50 por página
    page_number = request.GET.get('page')
    families_page = paginator.get_page(page_number)

    context = {
        'families': families_page,  # objeto de página
        'query': query,
        'show_ativas': show_ativas,
        'total': total,
        'total_ativas': total_ativas,
    }
    return render(request, 'AppLSD/family_list.html', context)

@login_required
def family_edit(request, pk):
    family = get_object_or_404(Family, pk=pk)

    if request.method == 'POST':
        form = FamilyForm(request.POST, request.FILES, instance=family)

        if form.is_valid():
            family = form.save(commit=False)
            family.editado_por = request.user
            family.save()

            registrar_log(
                request.user,
                family,
                'familia_editada',
                f'Família {family.responsible_name} atualizada.'
            )

            # Campo de arquivo único (file_info) OU múltiplos (documentos)
            # 1) Se você mantiver um único input:
            arquivo_atual = request.FILES.get('file_info')
            if arquivo_atual:
                DocumentoFamilia.objects.create(
                    family=family,
                    ano=timezone.now().year,
                    tipo='geral',
                    arquivo=arquivo_atual,
                )

            # 2) Se usar <input type="file" name="documentos" multiple>:
           

            return redirect('family_detail', pk=family.pk)
    else:
        form = FamilyForm(instance=family)

    return render(
        request,
        'AppLSD/family_form.html',
        {'form': form, 'family': family},
    )

@login_required
def family_create(request):
    if request.method == 'POST':
        form = FamilyForm(request.POST, request.FILES or None)
        formset = AlunoInlineFormSet(request.POST, request.FILES or None)

        try:
            form_ok = form.is_valid()
            formset_ok = formset.is_valid()

            logger.warning("family_create start | user=%s | form_ok=%s | formset_ok=%s",
                           request.user, form_ok, formset_ok)

            if form_ok and formset_ok:
                with transaction.atomic():
                    logger.warning("step 1: form.save(commit=False)")
                    family = form.save(commit=False)

                    logger.warning("step 2: set criado_por")
                    family.criado_por = request.user

                    logger.warning("step 3: family.save()")
                    family.save()

                    logger.warning("step 4: set formset.instance")
                    formset.instance = family

                    logger.warning("step 5: formset.save()")
                    formset.save()

                    logger.warning("step 6: registrar_log()")
                    registrar_log(
                        request.user,
                        family,
                        'familia_criada',
                        f'Família {family.responsible_name} cadastrada.'
                    )

                messages.success(request, 'Família cadastrada com sucesso.')
                return redirect('family_detail', pk=family.pk)

            messages.error(request, 'Não foi possível salvar. Verifique os campos.')

        except Exception:
            logger.exception("Erro inesperado em family_create | user=%s", request.user)
            messages.error(request, 'Erro interno ao salvar a família.')

    else:
        form = FamilyForm()
        formset = AlunoInlineFormSet()

    return render(request, 'AppLSD/family_form.html', {'form': form, 'formset': formset})

@login_required
def family_update(request, pk):
    family = get_object_or_404(Family, pk=pk)
    if request.method == 'POST':
        # guarda o valor antigo de file_info antes de processar o form
        file_info_antigo = family.file_info
        form = FamilyForm(request.POST, request.FILES, instance=family)
        formset = AlunoInlineFormSet(request.POST, request.FILES, instance=family)
        if form.is_valid() and formset.is_valid():
            family = form.save(commit=False) 
            family.editado_por = request.user           # auditoria
            family.save()
            formset.save()

            print("FILES:", request.FILES)  # DEBUG
            # arquivo do documento atual (campo EXTRA no template)
            arquivo_atual = request.FILES.get('file_info')

            print("ARQUIVO_ATUAL:", arquivo_atual)  # DEBUG

            if arquivo_atual:
                ano_atual = timezone.now().year
                DocumentoFamilia.objects.filter(family=family, ano=ano_atual).delete()
                DocumentoFamilia.objects.create(
                    family=family,
                    ano=ano_atual,
                    arquivo=arquivo_atual,
                )
                # restaura o file_info antigo para continuar sendo o "documento anterior"
                family.file_info = file_info_antigo
                family.save(update_fields=['file_info'])
            return redirect('family_detail', pk=family.pk)
    else:
        form = FamilyForm(instance=family)
        formset = AlunoInlineFormSet(instance=family)
    return render(request, 'families/family_form.html', {'form': form, 'formset': formset, 'family': family},)

def buscar_family(request):
    query = request.GET.get('query', '')
    families = Family.objects.filter(
        Q(responsible_name__icontains=query) | Q(registration_number__icontains=query) | Q(cpf__icontains=query)
    )[:10]
    results = [{
        'id': f.pk,
        'responsible_name': f.responsible_name,
        'cpf': f.cpf,
        'registration_number': f.registration_number
    } for f in families]
    return JsonResponse(results, safe=False)

# Gerar PDF do termo de responsabilidade da família
def family_term_pdf(request, pk):
    family = get_object_or_404(Family, pk=pk)
    today = date.today()

    context = {
        'family': family,
        'city': 'Maceió',
        'today': today,
        'day': today.day,
        'month': today.strftime('%B'),   # ou em português manualmente
        'year': today.year,
    }

    html_string = render_to_string('terms/family_term.html', context)
    html = HTML(string=html_string, base_url=request.build_absolute_uri())
    pdf_file = html.write_pdf()

    response = HttpResponse(pdf_file, content_type='application/pdf')
    # se quiser forçar download, use attachment:
    response['Content-Disposition'] = f'filename="termo_familia_{family.id}.pdf"'

    return response

@login_required
def adult_list(request):
    # Captura o termo de busca
    query = request.GET.get('q', '').strip()
    
    # Base queryset - ajuste o nome do modelo conforme seu código
    adultos = Adult.objects.select_related('family').all()  # ou Person, conforme seu modelo
    
    # Aplica filtro se houver busca
    if query:
        adultos = adultos.filter(
            Q(name__icontains=query) |
            Q(cpf__icontains=query) |
            Q(family__registration_number__icontains=query) |
            Q(family__responsible_name__icontains=query)
        )
    
    # Ordena por nome
    adultos = adultos.order_by('name')
    
    # Conta total
    total = adultos.count()
    
    # Paginação
    paginator = Paginator(adultos, 50)
    page_number = request.GET.get('page', 1)
    page_obj = paginator.get_page(page_number)
    
    context = {
        'adultos': page_obj,
        'query': query,
        'total': total,
    }
    
    return render(request, 'AppLSD/adult_list.html', context)

@login_required
def adult_create(request):
    family_id = request.GET.get('family_id')
    if not family_id:
        # Redirecione para a lista de famílias ou exiba uma mensagem de erro amigável
        return redirect('family_list')

    family = Family.objects.get(pk=family_id)
    
    if request.method == 'POST':
        form = AdultForm(request.POST)
        if form.is_valid():
            adult = form.save(commit=False)
            adult.family = family
            adult.criado_por = request.user             # auditoria
            adult.save()
            registrar_log(
                request.user,
                adult,
                'adulto_criado',
                f'Adulto {adult.name} incluído na família {adult.family.responsible_name}.'
            )
            return redirect('family_detail', pk=family.pk)
    else:
        form = AdultForm()
        form.fields['family'].initial = family.pk  # mantém selecionado
        form.fields['family'].widget = forms.HiddenInput()   # campo bloqueado
    return render(request, 'AppLSD/adult_form.html', {'form': form, 'family': family})

@login_required
def adult_detail(request, pk):
    adult = get_object_or_404(Adult, pk=pk)

    from django.contrib.contenttypes.models import ContentType
    ct = ContentType.objects.get_for_model(Adult)
    logs = AppLog.objects.filter(content_type=ct, object_id=adult.pk)
    
    context = {
        'adult': adult,
        'logs': logs,
    }
    return render(request, 'AppLSD/adult_detail.html', context)

@login_required
def adult_edit(request, pk):
    adult = get_object_or_404(Adult, pk=pk)
    if request.method == 'POST':
        form = AdultForm(request.POST, instance=adult)
        if form.is_valid():
            adult = form.save(commit=False)
            adult.editado_por = request.user            # auditoria
            adult.save()
            registrar_log(
                request.user,
                adult,
                'adulto_editado',
                f'Dados do adulto {adult.name} atualizados.'
            )
            return redirect('family_detail', pk=adult.family.pk)
    else:
        form = AdultForm(instance=adult)
    return render(request, 'AppLSD/adult_form.html', {'form': form, 'family': adult.family, 'adult': adult})

@login_required
def adult_delete_confirm(request, pk):
    adult = get_object_or_404(Adult, pk=pk)
    family_pk = adult.family.pk
    error = None  # Inicializa a variável error

    if request.method == 'POST':
        senha = request.POST.get('senha')
        usuario = request.user
        user_autenticado = authenticate(username=usuario.username, password=senha)
        if user_autenticado is not None:
            adult.delete()
            return redirect('family_detail', pk=family_pk)  # Redireciona para detalhes da família
        else:
            error = "Senha incorreta."

    return render(request, 'AppLSD/adult_confirm_delete.html', {
        'adult': adult,
        'error': error,
    })


def family_autocomplete(request):
    term = request.GET.get("q", "")
    familias = Family.objects.filter(responsible_name__icontains=term).order_by("responsible_name")[:20]
    results = [{"id": f.id, "text": f.responsible_name} for f in familias]
    return JsonResponse({"results": results})




@login_required
def aluno_list(request):
    # Captura o termo de busca
    query = request.GET.get('q', '').strip()
    show_frequentando = request.GET.get('show_frequentando')
    
    print('DEBUG show_frequentando =', show_frequentando)  # aqui

    # Base queryset
    alunos_qs = Aluno.objects.select_related('family', 'turma').all()
    # Filtra os alunos
    if query:
        alunos_qs = alunos_qs.filter(
            Q(name__icontains=query) |
            Q(cpf__icontains=query) |
            Q(family__registration_number__icontains=query) |
            Q(family__responsible_name__icontains=query)  # Busca pelo responsável
        )
    
    if show_frequentando == 'on':
        alunos_qs = alunos_qs.filter(status_lsd='Frequentando')  # ajuste ao valor real
    
        print(show_frequentando)
    # Ordena
    alunos_qs = alunos_qs.order_by('name')
    # Conta total ANTES da paginação
    total = alunos_qs.count()
    total_ativos = alunos_qs.filter(status_lsd='Frequentando').count()

      # Paginação
    paginator = Paginator(alunos_qs, 50)
    page_number = request.GET.get('page', 1)
    page_obj = paginator.get_page(page_number)
    
    context = {
        'alunos': page_obj,
        'query': query,
        'show_frequentando': show_frequentando,
        'total': total,
        'total_ativos': total_ativos,
    }
    
    return render(request, 'AppLSD/aluno_list.html', context)

@login_required
def aluno_create(request):
    family_id = request.GET.get('family_id')
    family = get_object_or_404(Family, pk=family_id) if family_id else None
    #if not family_id:
        # Redirecione para a lista de famílias ou exiba uma mensagem de erro amigável
     #   return redirect('family_list')

    #family = Family.objects.get(pk=family_id)
    if request.method == 'POST':
        form = AlunoForm(request.POST)
        if form.is_valid():
            try:
                aluno = form.save(commit=False)
                aluno.criado_por = request.user # registra quem criou
                # Associa família caso venha selecionada/fixa
                if family:
                    aluno.family = family
                aluno.save()
                if aluno.family:
                    family_name = aluno.family.responsible_name
                    registrar_log(
                        request.user,
                        aluno,
                        'aluno_criado',
                        f'Aluno {aluno.name} criado e vinculado à família {aluno.family.responsible_name}.'
                    )
                    logger.info(f"Aluno {aluno.name} (ID: {aluno.pk}) criado e vinculado à família {family_name}")
                    messages.success(request, f'Aluno {aluno.name} cadastrado com sucesso!')
            #return redirect('family_detail')  # Ou outro local desejado
                    return redirect('family_detail', pk=aluno.family.pk)
                else:
                    logger.info(f"Aluno {aluno.name} (ID: {aluno.pk}) criado sem família")
                    messages.success(request, f'Aluno {aluno.name} cadastrado com sucesso!')

                    return redirect('aluno_list')  # ajuste para sua view de lista de alunos
                    
            except Exception as e:
                logger.error(f"Erro ao criar aluno: {str(e)}")
                messages.error(request, f'Erro ao cadastrar aluno: {str(e)}')
        else:
            # Exibe erros do formulário
            logger.warning(f"Formulário inválido ao criar aluno: {form.errors}")
            messages.error(request, 'Por favor, corrija os erros no formulário.')
    else:
        
        # Se veio a família, preenche campo oculto. Senão, deixa campo visível
        initial_data = {'family': family} if family else {}
        form = AlunoForm(initial=initial_data)
        #form = AlunoForm()
        if family:
            #form.fields['family'].initial = family.pk  # mantém selecionado
            form.fields['family'].widget = forms.HiddenInput()   # campo bloqueado
    return render(request, 'AppLSD/aluno_form.html', {'form': form, 'family': family})

@login_required
def aluno_detail(request, pk):
    aluno = get_object_or_404(Aluno, pk=pk)

    educadora_status = request.user.groups.filter(name='Educadora').exists()

    turma_atual = aluno.turma

    ct = ContentType.objects.get_for_model(Aluno)
    logs = AppLog.objects.filter(content_type=ct, object_id=aluno.pk)

    historico_turmas = MovimentacaoTurmaAluno.objects.filter(
        aluno=aluno
    ).order_by('-data')

    pode_desligar_aluno = (
        not educadora_status and aluno.status_lsd != 'desligado'
    )

    is_educadora_da_turma = (
        educadora_status
        and aluno.turma is not None
        and aluno.turma.educadora == request.user
    )

    pode_mover_aluno = (
        request.user.is_superuser
        or is_coordenacao(request.user)
    )

    pode_adicionar_atividade = (
        request.user.is_superuser
        or is_coordenacao(request.user)
        or educadora_status
    )

    atividades_disponiveis = get_atividades_disponiveis_para_aluno(aluno)

    context = {
        'aluno': aluno,
        'logs': logs,
        'historico_turmas': historico_turmas,
        'is_educadora': educadora_status,
        'turma_atual': turma_atual,
        'pode_desligar_aluno': pode_desligar_aluno,
        'pode_mover_aluno': pode_mover_aluno,
        'pode_adicionar_atividade': pode_adicionar_atividade,
        'atividades_disponiveis': atividades_disponiveis,
    }
    return render(request, 'AppLSD/aluno_detail.html', context)

@login_required
def aluno_edit(request, pk):
    if usuario_tem_perfil(request.user, "colaborador"):
        raise PermissionDenied("Colaborador não pode editar aluno.")

    aluno = get_object_or_404(Aluno, pk=pk)
    status_anterior = aluno.status_lsd

    if request.method == 'POST':
        form = AlunoForm(request.POST, instance=aluno)
        if form.is_valid():
            aluno_editado = form.save(commit=False)
            novo_status = aluno_editado.status_lsd

            with transaction.atomic():
                if status_anterior != 'desligado' and novo_status == 'desligado':
                    motivo = 'Desligamento realizado pela edição do aluno'

                    atividades = list(aluno.atividades.all())

                    for atividade in atividades:
                        atividade.alunos.remove(aluno)

                        registrar_log(
                            request.user,
                            atividade,
                            'aluno_removido_desligamento',
                            f'Aluno {aluno.name} removido da atividade "{atividade.atividade}" por desligamento via edição. Motivo: {motivo}.'
                        )

                    if atividades:
                        registrar_log(
                            request.user,
                            aluno,
                            'aluno_removido_atividades',
                            f'Aluno removido de todas as atividades por desligamento via edição. Motivo: {motivo}.'
                        )

                    if aluno.turma:
                        turma_atual = aluno.turma
                        aluno.turma = None

                        registrar_log(
                            request.user,
                            turma_atual,
                            'aluno_removido_turma',
                            f'Aluno {aluno.name} removido da turma "{turma_atual.grupo}" por desligamento via edição. Motivo: {motivo}.'
                        )

                        registrar_log(
                            request.user,
                            aluno,
                            'aluno_removido_turma',
                            f'Aluno removido da turma "{turma_atual.grupo}" por desligamento via edição. Motivo: {motivo}.'
                        )

                    aluno.status_lsd = 'desligado'
                    aluno.editado_por = request.user
                    aluno.motivo_desligamento = motivo
                    aluno.data_desligamento = timezone.now()
                    aluno.desligado_por = request.user
                    aluno.save()

                    registrar_log(
                        request.user,
                        aluno,
                        'aluno_desligado',
                        f'Aluno {aluno.name} desligado com sucesso via edição. Motivo: {motivo}.'
                    )
                else:
                    aluno_editado.editado_por = request.user
                    aluno_editado.save()
                    form.save_m2m()

                    registrar_log(
                        request.user,
                        aluno_editado,
                        'aluno_editado',
                        f'Aluno {aluno_editado.name} editado com sucesso.'
                    )

            messages.success(request, 'Aluno atualizado com sucesso.')
            return redirect('aluno_detail', pk=aluno.pk)
    else:
        form = AlunoForm(instance=aluno)

    context = {
        'form': form,
        'aluno': aluno,
    }

    logger.debug(f"Ensino: {aluno.ensino}, Série: {aluno.serie}")
    return render(request, 'AppLSD/aluno_form.html', context)

@login_required
def aluno_delete_confirm(request, pk):
    # bloqueia colaboradores
    if usuario_tem_perfil(request.user, "colaborador"):
        raise PermissionDenied("Colaborador não pode excluir aluno.")
    aluno = get_object_or_404(Aluno, pk=pk)
    error = None

    if request.method == 'POST':
        senha = request.POST.get('senha')
        usuario = request.user
        user_autenticado = authenticate(username=usuario.username, password=senha)
        if user_autenticado is not None:
            # se quiser manter delete “real”, só delete()
            # se quiser já usar auditoria no futuro: trocar por soft_delete
            # aluno.soft_delete(request.user)
            aluno.delete()
            return redirect('aluno_list')
        else:
            error = "Senha incorreta."

    return render(request, 'AppLSD/aluno_confirm_delete.html', {
        'aluno': aluno,
        'error': error,
    })

def desligar_aluno_logic(*, request_user, aluno, motivo, origem='manual', familia=None):
    motivo = (motivo or '').strip() or 'Não informado'

    atividades = Activity.objects.filter(alunos=aluno)

    for atividade in atividades:
        atividade.alunos.remove(aluno)

        registrar_log(
            request_user,
            atividade,
            'aluno_removido_desligamento',
            f'Aluno {aluno.name} removido da atividade "{atividade.atividade}" por desligamento. Motivo: {motivo}.'
        )

    if atividades.exists():
        registrar_log(
            request_user,
            aluno,
            'aluno_removido_atividades',
            f'Aluno removido de todas as atividades por desligamento. Motivo: {motivo}.'
        )

    if aluno.turma:
        turma_atual = aluno.turma
        aluno.turma = None

        registrar_log(
            request_user,
            turma_atual,
            'aluno_removido_turma',
            f'Aluno {aluno.name} removido da turma "{turma_atual.grupo}" por desligamento. Motivo: {motivo}.'
        )

        registrar_log(
            request_user,
            aluno,
            'aluno_removido_turma',
            f'Aluno removido da turma "{turma_atual.grupo}" por desligamento. Motivo: {motivo}.'
        )

    aluno.status_lsd = 'Desligado'
    aluno.editado_por = request_user

    if hasattr(aluno, 'motivo_desligamento'):
        aluno.motivo_desligamento = motivo

    if hasattr(aluno, 'data_desligamento'):
        aluno.data_desligamento = timezone.now()

    if hasattr(aluno, 'desligado_por'):
        aluno.desligado_por = request_user

    aluno.save()

    registrar_log(
        request_user,
        aluno,
        'aluno_desligado',
        f'Aluno {aluno.name} desligado com sucesso. Motivo: {motivo}.'
    )

@login_required
def desativar_aluno(request, aluno_id):
    aluno = get_object_or_404(Aluno, id=aluno_id)

    if not (request.user.is_superuser or is_coordenacao(request.user)):
        return HttpResponseForbidden('Sem permissão')

    if request.method == 'POST':
        senha = request.POST.get('senha', '').strip()
        motivo = request.POST.get('motivo', '').strip()

        if not senha:
            messages.error(request, 'Senha obrigatória!')
            return redirect('aluno_detail', pk=aluno.id)

        if not check_password(senha, request.user.password):
            messages.error(request, 'Senha incorreta!')
            return redirect('aluno_detail', pk=aluno.id)

        if not motivo:
            messages.error(request, 'O motivo do desligamento é obrigatório.')
            return redirect('aluno_detail', pk=aluno.id)

        with transaction.atomic():
            desligar_aluno_logic(
                request_user=request.user,
                aluno=aluno,
                motivo=motivo,
                origem='manual'
            )

        messages.success(
            request,
            f'Aluno {aluno.name} desligado com sucesso, removido da turma e das atividades.'
        )
        return redirect('aluno_detail', pk=aluno.id)

    return render(request, 'AppLSD/desativar_aluno.html', {'aluno': aluno})


@login_required
def desativar_familia(request, familia_id):
    """
    Desativa uma família e todos os seus assistidos.
    Remove alunos de turmas e atividades, registrando no histórico.
    """
    familia = get_object_or_404(Family, id=familia_id)

    if request.method == 'POST':
        with transaction.atomic():
            alunos = Aluno.objects.filter(family=familia)

            for aluno in alunos:
                desligar_aluno_logic(
                    request_user=request.user,
                    aluno=aluno,
                    motivo=f'Desativação da família {familia.name}',
                    origem='familia_desativada',
                    familia=familia
                )

            familia.ativo = False
            familia.save()

            messages.success(
                request,
                f'Família {familia.name} desativada. '
                f'{alunos.count()} aluno(s) removido(s) de turmas e atividades.'
            )

            return redirect('familia_list')

    return render(request, 'desativar_familia.html', {'familia': familia})


def _remover_aluno_da_atividade_logic(request, aluno, atividade):
    """
    Lógica comum para remover aluno de atividade com validação de senha.
    Retorna True em caso de sucesso e False em caso de falha.
    """

    pode_remover = (
        request.user.is_superuser
        or is_coordenacao(request.user)
        or atividade.facilitador == request.user
        or (aluno.turma and aluno.turma.educadora == request.user)
    )

    if not pode_remover:
        messages.error(request, 'Sem permissão para remover este aluno da atividade.')
        return False

    if request.method != 'POST':
        messages.error(request, 'Método inválido para remoção.')
        return False

    senha = request.POST.get('senha', '').strip()
    if not senha:
        messages.error(request, 'Senha obrigatória!')
        return False

    if not check_password(senha, request.user.password):
        messages.error(request, 'Senha incorreta!')
        return False

    motivo = request.POST.get('motivo', '').strip()
    if not motivo:
        messages.error(request, 'O motivo da remoção é obrigatório.')
        return False

    if not aluno.atividades.filter(id=atividade.id).exists():
        messages.warning(request, 'O aluno já não está vinculado a esta atividade.')
        return False

    aluno.atividades.remove(atividade)

    registrar_log(
        request.user,
        atividade,
        'aluno_removido_da_atividade',
        f'Aluno {aluno.name} removido da atividade "{atividade.atividade}". Motivo: {motivo}.'
    )

    registrar_log(
        request.user,
        aluno,
        'aluno_removido_atividade',
        f'Aluno removido da atividade "{atividade.atividade}". Motivo: {motivo}.'
    )

    messages.success(request, 'Aluno removido da atividade.')
    return True

@login_required
def remover_aluno_da_atividade(request, activity_id, aluno_id):
    """
    Remove aluno da atividade - chamada da tela de DETALHES DA ATIVIDADE
    """
    atividade = get_object_or_404(Activity, id=activity_id)
    aluno = get_object_or_404(Aluno, id=aluno_id)

    _remover_aluno_da_atividade_logic(request, aluno, atividade)
    return redirect('activity_detail', activity_id=atividade.id)

@login_required
def remover_aluno_da_atividade_por_aluno(request, aluno_id, activity_id):
    """
    Remove aluno da atividade - chamada da tela de DETALHES DO ALUNO
    URL: /alunos/<aluno_id>/remover-da-atividade/<activity_id>/
    """
    aluno = get_object_or_404(Aluno, id=aluno_id)
    atividade = get_object_or_404(Activity, id=activity_id)

    if _remover_aluno_da_atividade_logic(request, aluno, atividade):
        return redirect('aluno_detail', pk=aluno.id)
    
    return redirect('aluno_detail', pk=aluno.id)


@login_required
def turma_detail(request, turma_id):
    turma = get_object_or_404(Turma, id=turma_id)
    alunos = Aluno.objects.filter(turma=turma).order_by('name')
    total_alunos_turma = alunos.count()

    coordenacao_status = is_coordenacao(request.user)
    educadora_status = is_educadora(request.user)

    is_educadora_desta_turma = (
        educadora_status and turma.educadora == request.user
    )

    today = timezone.localdate()

    data_param = request.GET.get('data')
    if data_param:
        try:
            data_referencia = datetime.strptime(data_param, '%Y-%m-%d').date()
        except ValueError:
            data_referencia = today
    else:
        data_referencia = today

    aniversariantes_mes = alunos.filter(
        birth_date__month=today.month
    ).order_by('birth_date', 'name')

    aniversariantes_dia = aniversariantes_mes.filter(birth_date__day=today.day)

    frequencia_referencia = FrequenciaTurma.objects.filter(
        turma=turma,
        data=data_referencia
    ).first()

    pode_editar = False
    pode_visualizar = False
    pode_iniciar = False

    if frequencia_referencia:
        pode_visualizar = coordenacao_status or is_educadora_desta_turma
        pode_editar = pode_editar_frequencia_turma(request.user, turma)
    else:
        if coordenacao_status or is_educadora_desta_turma:
            pode_iniciar = True
            pode_visualizar = True

    ct = ContentType.objects.get_for_model(Turma)
    logs = AppLog.objects.filter(
        content_type=ct,
        object_id=turma.id
    ).order_by('-criado_em')

    context = {
        'turma': turma,
        'logs': logs,
        'alunos': alunos,
        'frequencia_hoje': FrequenciaTurma.objects.filter(turma=turma, data=today).first(),
        'frequencia_referencia': frequencia_referencia,
        'ja_tem_frequencia': frequencia_referencia is not None,
        'pode_editar': pode_editar,
        'pode_visualizar': pode_visualizar,
        'pode_iniciar': pode_iniciar,
        'is_coordenacao': coordenacao_status,
        'is_educadora': educadora_status,
        'is_educadora_desta_turma': is_educadora_desta_turma,
        'total_alunos_turma': total_alunos_turma,
        'pode_adicionar_alunos': coordenacao_status,
        'pode_editar_alunos': coordenacao_status,
        'pode_mover_alunos': coordenacao_status or is_educadora_desta_turma,
        'aniversariantes_mes': aniversariantes_mes,
        'aniversariantes_dia': aniversariantes_dia,
        'data_referencia': data_referencia,
        'data_hoje': today,
        'consultando_hoje': data_referencia == today,
    }
    return render(request, 'AppLSD/turma_detail.html', context)


def get_tipo_perfil(user):
    """
    Retorna o tipo de perfil do usuário
    """
    if user.is_superuser:
        return 'admin'
    
    perfis = user.perfis.all()
    if perfis.exists():
        return perfis.first().tipo_perfil
    
    return None

"""
def is_coordenacao(user):
    
    #Verifica se o usuário é coordenação
        tipo_perfil = get_tipo_perfil(user)
    return tipo_perfil == 'coordenacao' or user.is_superuser
"""

def is_educadora(user):
    """
    Verifica se o usuário é educadora   
    """
    tipo_perfil = get_tipo_perfil(user)
    return tipo_perfil == 'educadora'


class TurmaListView(ListView):
    model = Turma
    template_name = "AppLSD/turma_list.html"   # use o seu template atual
    context_object_name = "page_obj"           # para reaproveitar o template
    paginate_by = 20                          # mesmo valor do Paginator

    def get_queryset(self):
        qs = Turma.objects.all().order_by('faixa_etaria','educadora', 'grupo', 'turno')

        turno = self.request.GET.get("turno")
        ano = self.request.GET.get("ano")
        faixa = self.request.GET.get("faixa_etaria")
        educadora = self.request.GET.get("educadora")

        if turno:
            qs = qs.filter(turno=turno)
        if ano:
            qs = qs.filter(ano_letivo=ano)
        if faixa:
            qs = qs.filter(faixa_etaria=faixa)
        if educadora:
            qs = qs.filter(educadora__first_name__icontains=educadora)


        return qs

@login_required
def turma_create(request):
    if request.method == 'POST':
        form = TurmaForm(request.POST)
        if form.is_valid():
            turma = form.save(commit=False)
            turma.criado_por = request.user              # auditoria
            form.save()
            registrar_log(
                request.user,
                turma,
                'turma_criada',
                f'Turma {turma.grupo} criada para o turno {turma.turno}.'
            )
            return redirect('turma_list')
    else:
        form = TurmaForm()
    return render(request, 'AppLSD/turma_form.html', {'form': form})

@login_required
def turma_edit(request, pk):
    turma = get_object_or_404(Turma, pk=pk)
    if request.method == 'POST':
        form = TurmaForm(request.POST, instance=turma)
        if form.is_valid():
            turma = form.save(commit=False)
            turma.editado_por = request.user            # auditoria
            form.save()
            registrar_log(
                request.user,
                turma,
                'turma_editada',
                f'Turma {turma.grupo} editada para o turno {turma.turno}.'
            )
            return redirect('turma_list')
    else:
        form = TurmaForm(instance=turma)
    return render(request, 'AppLSD/turma_form.html', {'form': form})

@login_required
def turma_delete(request, pk):
    turma = get_object_or_404(Turma, pk=pk)
    if request.method == 'POST':
        #turma.soft_delete(request.user)               # Se quiser deletar turma usar soft delte
        turma.delete()
        return redirect('turma_list')
    return render(request, 'AppLSD/turma_confirm_delete.html', {'turma': turma})

@login_required
def iniciar_frequencia_turma(request, turma_id):
    """
    Inicia uma nova frequência para a turma.
    Só cria/salva frequência no POST.
    Permite iniciar frequência em uma data informada.
    Também permite registrar quando não houve aula.
    """
    turma = get_object_or_404(Turma, id=turma_id)

    if not (is_coordenacao(request.user) or (is_educadora(request.user) and turma.educadora == request.user)):
        raise PermissionDenied("Você não tem permissão para iniciar a frequência desta turma.")

    data_param = request.GET.get('data') or request.POST.get('data')

    if data_param:
        try:
            data_referencia = datetime.strptime(data_param, '%Y-%m-%d').date()
        except ValueError:
            data_referencia = timezone.now().date()
    else:
        data_referencia = timezone.now().date()

    alunos = Aluno.objects.filter(turma=turma).select_related('family').order_by('name')

    status_aula = request.POST.get(
        'status_aula',
        FrequenciaTurma.StatusAula.NORMAL
    ) if request.method == 'POST' else FrequenciaTurma.StatusAula.NORMAL

    motivo_nao_aula = request.POST.get('motivo_nao_aula', '').strip() if request.method == 'POST' else ''
    observacao_nao_aula = request.POST.get('observacao_nao_aula', '').strip() if request.method == 'POST' else ''

    presencas_form = {
        aluno.id: {'presente': True, 'motivo_falta': ''}
        for aluno in alunos
    }
    erros_motivo = {}

    if request.method == "POST":
        erros = []

        for aluno in alunos:
            presente = f'presente_{aluno.id}' in request.POST
            motivo_falta = request.POST.get(f'motivo_{aluno.id}', '').strip()

            presencas_form[aluno.id] = {
                'presente': presente,
                'motivo_falta': motivo_falta,
            }

        if status_aula == FrequenciaTurma.StatusAula.NAO_HOUVE:
            if not motivo_nao_aula:
                messages.error(request, 'Selecione o motivo de não haver aula.')

                return render(
                    request,
                    'AppLSD/frequencia_turma_iniciar.html',
                    {
                        'turma': turma,
                        'alunos': alunos,
                        'data_hoje': data_referencia,
                        'status_aula': status_aula,
                        'motivo_nao_aula': motivo_nao_aula,
                        'observacao_nao_aula': observacao_nao_aula,
                        'motivos_nao_aula': FrequenciaTurma.MotivoNaoAula.choices,
                        'presencas_form': presencas_form,
                        'erros_motivo': erros_motivo,
                    }
                )
        else:
            for aluno in alunos:
                presente = presencas_form[aluno.id]['presente']
                motivo_falta = presencas_form[aluno.id]['motivo_falta']

                if not presente and not motivo_falta:
                    erros.append(f'Selecione o motivo da falta para o aluno {aluno.name}.')
                    erros_motivo[aluno.id] = 'Selecione um motivo da falta.'

            if erros:
                messages.error(request, 'Existem alunos com falta sem motivo selecionado.')

                return render(
                    request,
                    'AppLSD/frequencia_turma_iniciar.html',
                    {
                        'turma': turma,
                        'alunos': alunos,
                        'data_hoje': data_referencia,
                        'status_aula': status_aula,
                        'motivo_nao_aula': motivo_nao_aula,
                        'observacao_nao_aula': observacao_nao_aula,
                        'motivos_nao_aula': FrequenciaTurma.MotivoNaoAula.choices,
                        'presencas_form': presencas_form,
                        'erros_motivo': erros_motivo,
                    }
                )

        with transaction.atomic():
            chamada, created = FrequenciaTurma.objects.get_or_create(
                turma=turma,
                data=data_referencia,
                defaults={
                    'presente': True,
                    'criado_por': request.user,
                }
            )

            chamada.status_aula = status_aula
            chamada.motivo_nao_aula = motivo_nao_aula if status_aula == FrequenciaTurma.StatusAula.NAO_HOUVE else ''
            chamada.observacao_nao_aula = observacao_nao_aula if status_aula == FrequenciaTurma.StatusAula.NAO_HOUVE else ''

            if created:
                chamada.criado_por = request.user
            else:
                chamada.editado_por = request.user

            if status_aula == FrequenciaTurma.StatusAula.NAO_HOUVE:
                chamada.presente = False
                chamada.save()

                FrequenciaAluno.objects.filter(chamada=chamada).delete()

                messages.success(request, 'Registro salvo: não houve aula nesta data.')

                registrar_log(
                    request.user,
                    turma,
                    'frequencia_nao_houve_aula',
                    f'Na turma {turma.grupo}, foi registrado que não houve aula em {data_referencia.strftime("%d/%m/%Y")} '
                    f'por motivo de {chamada.get_motivo_nao_aula_display()}.',
                )

                return redirect('frequencia_visualizar', frequencia_id=chamada.id)

            chamada.presente = True
            chamada.save()

            for aluno in alunos:
                presente = presencas_form[aluno.id]['presente']
                motivo_falta = presencas_form[aluno.id]['motivo_falta']

                FrequenciaAluno.objects.update_or_create(
                    chamada=chamada,
                    aluno=aluno,
                    defaults={
                        'presente': presente,
                        'status': FrequenciaAluno.StatusPresenca.PRESENTE if presente else FrequenciaAluno.StatusPresenca.FALTA,
                        'motivo_falta': motivo_falta if not presente else '',
                    }
                )

        messages.success(request, 'Frequência registrada com sucesso!')

        registrar_log(
            request.user,
            turma,
            'frequencia_criada',
            f'Frequência do dia {data_referencia.strftime("%d/%m/%Y")} registrada pela educadora '
            f'{request.user.get_full_name() or request.user.username}.',
        )
        return redirect('frequencia_visualizar', frequencia_id=chamada.id)

    return render(
        request,
        'AppLSD/frequencia_turma_iniciar.html',
        {
            'turma': turma,
            'alunos': alunos,
            'data_hoje': data_referencia,
            'status_aula': FrequenciaTurma.StatusAula.NORMAL,
            'motivo_nao_aula': '',
            'observacao_nao_aula': '',
            'motivos_nao_aula': FrequenciaTurma.MotivoNaoAula.choices,
            'presencas_form': presencas_form,
            'erros_motivo': erros_motivo,
        }
    )


@login_required
def visualizar_frequencia_turma(request, frequencia_id):
    chamada_base = get_object_or_404(
        FrequenciaTurma.objects.select_related('turma', 'turma__educadora'),
        id=frequencia_id
    )

    turma = chamada_base.turma

    if not (
        is_coordenacao(request.user)
        or (is_educadora(request.user) and turma.educadora == request.user)
    ):
        raise PermissionDenied("Você não tem permissão para visualizar esta frequência.")

    data_param = request.GET.get('data')

    if data_param:
        try:
            data_consulta = datetime.strptime(data_param, '%Y-%m-%d').date()
        except ValueError:
            data_consulta = chamada_base.data
    else:
        data_consulta = chamada_base.data

    chamada = FrequenciaTurma.objects.filter(
        turma=turma,
        data=data_consulta
    ).select_related('turma', 'turma__educadora').first()

    alunos_turma = Aluno.objects.filter(
        turma=turma
    ).select_related('family').order_by('name')

    chamada_nao_encontrada = chamada is None

    alunos_com_presenca = []
    presentes = 0
    faltas = 0
    total_alunos = alunos_turma.count()

    if chamada_nao_encontrada:
        chamada = chamada_base
        presencas = FrequenciaAluno.objects.none()
    else:
        presencas = FrequenciaAluno.objects.filter(
            chamada=chamada
        ).select_related('aluno', 'aluno__family').order_by('aluno__name')

    presencas_map = {p.aluno_id: p for p in presencas}

    for aluno in alunos_turma:
        presenca = presencas_map.get(aluno.id)
        alunos_com_presenca.append({
            'aluno': aluno,
            'presenca': presenca,
        })

    if not chamada_nao_encontrada and chamada.status_aula != FrequenciaTurma.StatusAula.NAO_HOUVE:
        presentes = sum(
            1 for item in alunos_com_presenca
            if item['presenca'] and item['presenca'].presente
        )
        faltas = sum(
            1 for item in alunos_com_presenca
            if item['presenca'] and not item['presenca'].presente
        )

    pode_editar = pode_editar_frequencia_turma(request.user, turma)
    pode_iniciar_data = (
        is_coordenacao(request.user)
        or (is_educadora(request.user) and turma.educadora == request.user)
    )

    context = {
        'chamada': chamada,
        'turma': turma,
        'presencas': presencas,
        'alunos_turma': alunos_turma,
        'alunos_com_presenca': alunos_com_presenca,
        'presentes': presentes,
        'faltas': faltas,
        'total_alunos': total_alunos,
        'pode_editar': pode_editar,
        'chamada_nao_encontrada': chamada_nao_encontrada,
        'data_pesquisada': data_consulta.strftime('%Y-%m-%d') if data_param else '',
        'data_pesquisada_br': data_consulta.strftime('%d/%m/%Y') if data_param else '',
        'data_busca': data_consulta.strftime('%Y-%m-%d'),
        'pode_iniciar_data': pode_iniciar_data,
    }

    return render(request, 'AppLSD/frequencia_turma_visualizar.html', context)

#@user_passes_test(is_coordenacao)  # ✅ Apenas coordenação pode acessar
@login_required
def editar_frequencia_turma(request, frequencia_id):
    chamada = get_object_or_404(FrequenciaTurma, id=frequencia_id)

    if not pode_editar_frequencia_turma(request.user, chamada.turma):
        raise PermissionDenied("Você não tem permissão para editar esta frequência.")

    alunos_turma = Aluno.objects.filter(
        turma=chamada.turma
    ).select_related('family').order_by('name')

    presencas = FrequenciaAluno.objects.filter(
        chamada=chamada
    ).select_related('aluno').order_by('aluno__name')

    presencas_map = {p.aluno_id: p for p in presencas}

    presencas_form = {}
    for aluno in alunos_turma:
        presenca = presencas_map.get(aluno.id)
        presencas_form[aluno.id] = {
            'presente': presenca.presente if presenca else True,
            'motivo_falta': presenca.motivo_falta if presenca and not presenca.presente else '',
        }

    erros_motivo = {}

    if request.method == 'POST':
        status_aula = request.POST.get('status_aula', FrequenciaTurma.StatusAula.NORMAL)
        motivo_nao_aula = request.POST.get('motivo_nao_aula', '').strip()
        observacao_nao_aula = request.POST.get('observacao_nao_aula', '').strip()

        erros = {}

        for aluno in alunos_turma:
            presente = f'presente_{aluno.id}' in request.POST
            motivo_falta = request.POST.get(f'motivo_{aluno.id}', '').strip()

            presencas_form[aluno.id] = {
                'presente': presente,
                'motivo_falta': motivo_falta,
            }

        if status_aula == FrequenciaTurma.StatusAula.NAO_HOUVE:
            if not motivo_nao_aula:
                messages.error(request, 'Selecione o motivo de não haver aula.')
            else:
                with transaction.atomic():
                    chamada.status_aula = status_aula
                    chamada.motivo_nao_aula = motivo_nao_aula
                    chamada.observacao_nao_aula = observacao_nao_aula
                    chamada.editado_por = request.user
                    chamada.presente = False
                    chamada.save()

                    FrequenciaAluno.objects.filter(chamada=chamada).delete()

                messages.success(request, 'Registro atualizado: não houve aula nesta data.')

                registrar_log(
                    request.user,
                    chamada.turma,
                    'frequencia_editada_nao_houve_aula',
                    f'Frequência de {chamada.data.strftime("%d/%m/%Y")} alterada para "não houve aula" '
                    f'pela coordenadora ({request.user.get_full_name() or request.user.username}).',
                )
                return redirect('turma_detail', turma_id=chamada.turma.id)
        else:
            for aluno in alunos_turma:
                presente = presencas_form[aluno.id]['presente']
                motivo_falta = presencas_form[aluno.id]['motivo_falta']

                if not presente and not motivo_falta:
                    erros_motivo[aluno.id] = 'Selecione um motivo da falta.'

            if erros_motivo:
                messages.error(request, 'Existem alunos com falta sem motivo selecionado.')
            else:
                with transaction.atomic():
                    chamada.status_aula = FrequenciaTurma.StatusAula.NORMAL
                    chamada.motivo_nao_aula = ''
                    chamada.observacao_nao_aula = ''
                    chamada.editado_por = request.user
                    chamada.presente = True
                    chamada.save()

                    for aluno in alunos_turma:
                        presente = presencas_form[aluno.id]['presente']
                        motivo_falta = presencas_form[aluno.id]['motivo_falta']

                        FrequenciaAluno.objects.update_or_create(
                            chamada=chamada,
                            aluno=aluno,
                            defaults={
                                'presente': presente,
                                'status': FrequenciaAluno.StatusPresenca.PRESENTE if presente else FrequenciaAluno.StatusPresenca.FALTA,
                                'motivo_falta': motivo_falta if not presente else ''
                            }
                        )

                messages.success(request, 'Frequência atualizada com sucesso!')

                registrar_log(
                    request.user,
                    chamada.turma,
                    'frequencia_editada',
                    f'Frequência de {chamada.data.strftime("%d/%m/%Y")} editada pela coordenadora ({request.user.get_full_name() or request.user.username}).',
                )
                return redirect('turma_detail', turma_id=chamada.turma.id)

        chamada.status_aula = status_aula
        chamada.motivo_nao_aula = motivo_nao_aula if status_aula == FrequenciaTurma.StatusAula.NAO_HOUVE else ''
        chamada.observacao_nao_aula = observacao_nao_aula if status_aula == FrequenciaTurma.StatusAula.NAO_HOUVE else ''

    alunos_com_presenca = []
    for aluno in alunos_turma:
        alunos_com_presenca.append({
            'aluno': aluno,
            'presenca': presencas_form.get(aluno.id),
        })

    context = {
        'chamada': chamada,
        'alunos_com_presenca': alunos_com_presenca,
        'modo_edicao': True,
        'motivos_nao_aula': FrequenciaTurma.MotivoNaoAula.choices,
        'erros_motivo': erros_motivo,
    }

    return render(request, 'AppLSD/frequencia_turma_editar.html', context)

@login_required
def activity_detail(request, activity_id):
    atividade = get_object_or_404(Activity, id=activity_id)

    hoje = timezone.now().date()
    ja_tem_frequencia = FrequenciaAtividade.objects.filter(
        atividade=atividade,
        data=hoje,
    ).exists()

    alunos_atividade = Aluno.objects.filter(
        atividades=atividade
    ).order_by('name')

    ct = ContentType.objects.get_for_model(Activity)
    logs = AppLog.objects.filter(
        content_type=ct,
        object_id=atividade.pk
    ).order_by('-criado_em')

    user_is_coordenacao = is_coordenacao(request.user) or request.user.is_superuser
    user_is_educadora = is_educadora(request.user)

    context = {
        'atividade': atividade,
        'ja_tem_frequencia': ja_tem_frequencia,
        'tem_frequencia_hoje': ja_tem_frequencia,
        'logs': logs,
        'alunos_atividade': alunos_atividade,
        'is_coordenacao': user_is_coordenacao,
        'is_educadora': user_is_educadora,
        'pode_visualizar': user_is_coordenacao or user_is_educadora,
        'pode_iniciar': user_is_coordenacao or user_is_educadora,
        'pode_editar': user_is_coordenacao,
    }

    return render(request, 'AppLSD/activity_detail.html', context)


@login_required
def activity_list(request):
    atividades = Activity.objects.all().select_related('facilitador')

    atividade_nome = request.GET.get("atividade") or ""
    facilitador_id = request.GET.get("facilitador") or ""
    dia_semana = request.GET.get("dia_semana") or ""
    turno = request.GET.get("turno") or ""
    tipo = request.GET.get("tipo") or ""

    if atividade_nome:
        atividades = atividades.filter(atividade__icontains=atividade_nome)
    if facilitador_id:
        atividades = atividades.filter(facilitador_id=facilitador_id)
    if dia_semana:
        atividades = atividades.filter(dia_semana__icontains=dia_semana)
    if turno:
        atividades = atividades.filter(turno=turno)
    if tipo:
        atividades = atividades.filter(tipo=tipo)

    atividades = atividades.order_by('atividade')

    facilitadores = User.objects.filter(
        id__in=Activity.objects.values('facilitador_id')
    ).order_by('first_name').distinct()

    nomes_atividades = (
        Activity.objects.order_by('atividade')
        .values_list('atividade', flat=True)
        .distinct()
    )

    tipos_atividades = Activity.TIPO_CHOICES
    dias_semana_choices = Activity.DIAS_SEMANAS_CHOICES
    turnos = Activity.ESCOLHA_TURNO

    paginator = Paginator(atividades, 100)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    context = {
        'page_obj': page_obj,
        'facilitadores': facilitadores,
        'nomes_atividades': nomes_atividades,
        'tipos_atividades': tipos_atividades,
        'dias_semana_choices': dias_semana_choices,
        'turnos': turnos,
        'filtro_atividade': atividade_nome,
        'filtro_facilitador': facilitador_id,
        'filtro_dia_semana': dia_semana,
        'filtro_turno': turno,
        'filtro_tipo': tipo,
        'is_coordenacao': is_coordenacao(request.user),
        'is_educadora': is_educadora(request.user),
    }
    return render(request, 'AppLSD/activity_list.html', context)



@login_required
def activity_create(request):
    if request.method == 'POST':
        form = ActivityForm(request.POST)
        if form.is_valid():
            activity = form.save(commit=False)
            activity.criado_por = request.user          # auditoria
            activity.save()
            registrar_log(
                request.user,
                activity,
                'atividade_criada',
                f'Atividade {activity.atividade} criada para o turno {activity.turno}.'
            )
            form.save_m2m()
            return redirect('activity_list')
    else:
        form = ActivityForm()
    return render(request, 'AppLSD/activity_form.html', {'form': form})

def activity_edit(request, pk):
    atividade = get_object_or_404(Activity, pk=pk)
    if request.method == 'POST':
        form = ActivityForm(request.POST, instance=atividade)
        if form.is_valid():
            atividade = form.save(commit=False)
            atividade.editado_por = request.user        # auditoria
            atividade.save()
            form.save_m2m()
            registrar_log(
                request.user,
                atividade,
                'atividade_editada',
                f'Atividade {atividade.atividade} editada.'
            )
            
            return redirect('activity_list')
    else:
        form = ActivityForm(instance=atividade)
    return render(request, 'AppLSD/activity_form.html', {'form': form})

def activity_delete(request, pk):
    activity = get_object_or_404(Activity, pk=pk)
    if request.method == 'POST':
        #activity.soft_delete(request.user)              # se quiser soft delete
        # ou activity.delete() se ainda preferir exclusão real
        activity.delete()
        return redirect('activity_list')
    return render(request, 'AppLSD/activity_confirm_delete.html', {'activity': activity})


@login_required
def iniciar_frequencia_activity(request, activity_id):
    """
    Inicia ou registra a frequência de uma atividade.
    Educadoras registram a presença dos alunos vinculados à atividade.
    Coordenação pode registrar frequência de qualquer aluno da atividade.
    Se já houver falta registrada na turma na mesma data, replica para a atividade.
    """
    atividade = get_object_or_404(Activity, id=activity_id)

    user_is_coordenacao = request.user.is_superuser or is_coordenacao(request.user)

    if user_is_coordenacao:
        alunos = Aluno.objects.filter(atividades=atividade).select_related(
            'family', 'turma', 'turma__educadora'
        ).order_by('name')
    else:
        alunos = Aluno.objects.filter(
            atividades=atividade,
            turma__educadora=request.user,
        ).select_related(
            'family', 'turma', 'turma__educadora'
        ).order_by('name')

    data_param = request.POST.get('data') if request.method == 'POST' else request.GET.get('data')
    if data_param:
        try:
            data_chamada = date.fromisoformat(data_param)
        except ValueError:
            data_chamada = timezone.now().date()
    else:
        data_chamada = timezone.now().date()

    presencas_form = {}

    motivos_falta_choices = FrequenciaAtividade._meta.get_field('motivo_falta').choices
    motivos_validos = {choice[0] for choice in motivos_falta_choices}

    if request.method == 'POST':
        erros = []

        ja_existia = FrequenciaAtividade.objects.filter(
            atividade=atividade,
            data=data_chamada,
        ).exists()

        for aluno in alunos:
            presente = f'presente_{aluno.id}' in request.POST
            motivo_falta = request.POST.get(f'motivo_{aluno.id}', '').strip()

            presencas_form[aluno.id] = {
                'presente': presente,
                'motivo_falta': motivo_falta,
            }

            if not presente:
                if not motivo_falta:
                    erros.append(f'O aluno {aluno.name} está com falta e precisa ter um motivo selecionado.')
                elif motivo_falta not in motivos_validos:
                    erros.append(f'O motivo da falta do aluno {aluno.name} é inválido.')

        if erros:
            for erro in erros:
                messages.error(request, erro)

            return render(request, 'AppLSD/frequencia_activity_iniciar.html', {
                'atividade': atividade,
                'alunos': alunos,
                'data': data_chamada,
                'presencas_form': presencas_form,
                'motivos_falta': motivos_falta_choices,
            })

        for aluno in alunos:
            presente = f'presente_{aluno.id}' in request.POST
            motivo_falta = request.POST.get(f'motivo_{aluno.id}', '').strip()

            FrequenciaAtividade.objects.update_or_create(
                atividade=atividade,
                aluno=aluno,
                data=data_chamada,
                defaults={
                    'presente': presente,
                    'motivo_falta': motivo_falta if not presente else '',
                    'criado_por': request.user,
                }
            )

        messages.success(request, 'Frequência da atividade registrada com sucesso!')

        registrar_log(
            request.user,
            atividade,
            'frequencia_atividade_criada' if not ja_existia else 'frequencia_atividade_editada',
            (
                f'Frequência da atividade "{atividade.atividade}" em '
                f'{data_chamada.strftime("%d/%m/%Y")} registrada por '
                f'{request.user.get_full_name() or request.user.username}.'
            ),
        )

        return redirect('frequencia_activity_visualizar', activity_id=atividade.id)

    for aluno in alunos:
        freq_aluno_turma = FrequenciaAluno.objects.filter(
            aluno=aluno,
            chamada__data=data_chamada
        ).select_related('chamada').first()

        if freq_aluno_turma:
            presencas_form[aluno.id] = {
                'presente': freq_aluno_turma.presente,
                'motivo_falta': '' if freq_aluno_turma.presente else (freq_aluno_turma.motivo_falta or ''),
            }
        else:
            presencas_form[aluno.id] = {
                'presente': True,
                'motivo_falta': '',
            }

    return render(request, 'AppLSD/frequencia_activity_iniciar.html', {
        'atividade': atividade,
        'alunos': alunos,
        'data': data_chamada,
        'presencas_form': presencas_form,
        'motivos_falta': motivos_falta_choices,
    })

@login_required
def visualizar_frequencia_activity(request, activity_id):
    """
    Visualiza a frequência da atividade.
    Coordenação e educadoras podem visualizar.
    """
    atividade = get_object_or_404(Activity, id=activity_id)

    user_is_coordenacao = is_coordenacao(request.user) or request.user.is_superuser
    user_is_educadora = is_educadora(request.user)

    if not user_is_coordenacao and not user_is_educadora:
        messages.error(request, 'Você não tem permissão para visualizar a frequência desta atividade.')
        return redirect('activity_detail', activity_id=atividade.id)

    data = request.GET.get('data')
    if data:
        try:
            data_chamada = date.fromisoformat(data)
        except ValueError:
            messages.warning(request, 'Data inválida. Exibindo a data de hoje.')
            data_chamada = timezone.now().date()
    else:
        data_chamada = timezone.now().date()

    presencas = FrequenciaAtividade.objects.filter(
        atividade=atividade,
        data=data_chamada,
    ).select_related('aluno', 'aluno__turma', 'aluno__family').order_by('aluno__name')

    if user_is_educadora and not user_is_coordenacao:
        presencas = presencas.filter(aluno__turma__educadora=request.user)

    chamada = presencas.first()

    context = {
        'atividade': atividade,
        'data_chamada': data_chamada,
        'chamada': chamada,
        'presencas': presencas,
        'is_coordenacao': user_is_coordenacao,
        'is_educadora': user_is_educadora,
        'pode_editar': user_is_coordenacao and presencas.exists(),
        'total_presencas': presencas.count(),
    }
    return render(request, 'AppLSD/frequencia_activity_visualizar.html', context)


@user_passes_test(is_coordenacao)
def editar_frequencia_activity(request, activity_id):
    """
    Edita a frequência da atividade.
    Apenas coordenação pode editar.
    """
    atividade = get_object_or_404(Activity, id=activity_id)

    alunos_atividade = Aluno.objects.filter(
        atividades=atividade
    ).select_related(
        'family', 'turma', 'turma__educadora'
    ).order_by('name')

    data_param = request.POST.get('data') if request.method == 'POST' else request.GET.get('data')

    if data_param:
        try:
            data_chamada = date.fromisoformat(data_param)
        except ValueError:
            data_chamada = timezone.now().date()
    else:
        data_chamada = timezone.now().date()

    presencas_qs = FrequenciaAtividade.objects.filter(
        atividade=atividade,
        data=data_chamada
    ).select_related('aluno')

    presencas_dict = {p.aluno.id: p for p in presencas_qs}

    motivos_falta_choices = FrequenciaAtividade._meta.get_field('motivo_falta').choices
    motivos_validos = {choice[0] for choice in motivos_falta_choices}

    if request.method == 'POST':
        erros = []
        presencas_form = {}

        for aluno in alunos_atividade:
            presente = f'presente_{aluno.id}' in request.POST
            motivo_falta = request.POST.get(f'motivo_{aluno.id}', '').strip()

            presencas_form[aluno.id] = {
                'presente': presente,
                'motivo_falta': motivo_falta,
            }

            if not presente:
                if not motivo_falta:
                    erros.append(
                        f'O aluno {aluno.name} está com falta e precisa ter um motivo selecionado.'
                    )
                elif motivo_falta not in motivos_validos:
                    erros.append(
                        f'O motivo da falta do aluno {aluno.name} é inválido.'
                    )

        if erros:
            for erro in erros:
                messages.error(request, erro)

            context = {
                'atividade': atividade,
                'data_chamada': data_chamada,
                'presencas': presencas_form,
                'alunos_atividade': alunos_atividade,
                'modo_edicao': True,
                'motivos_falta': motivos_falta_choices,
            }
            return render(request, 'AppLSD/frequencia_activity_editar.html', context)

        for aluno in alunos_atividade:
            presente = f'presente_{aluno.id}' in request.POST
            motivo_falta = request.POST.get(f'motivo_{aluno.id}', '').strip()

            FrequenciaAtividade.objects.update_or_create(
                atividade=atividade,
                aluno=aluno,
                data=data_chamada,
                defaults={
                    'presente': presente,
                    'motivo_falta': motivo_falta if not presente else '',
                    'editado_por': request.user,
                }
            )

        messages.success(request, 'Frequência atualizada com sucesso!')

        registrar_log(
            request.user,
            atividade,
            'frequencia_atividade_editada',
            f'Frequência da atividade "{atividade.atividade}" em '
            f'{data_chamada.strftime("%d/%m/%Y")} editada pela coordenação '
            f'({request.user.get_full_name() or request.user.username}).',
        )

        return redirect('activity_detail', activity_id=atividade.id)

    presencas_form = {}
    for aluno in alunos_atividade:
        presenca = presencas_dict.get(aluno.id)
        presencas_form[aluno.id] = {
            'presente': presenca.presente if presenca else True,
            'motivo_falta': presenca.motivo_falta if presenca and not presenca.presente else '',
        }

    context = {
        'atividade': atividade,
        'data_chamada': data_chamada,
        'presencas': presencas_form,
        'alunos_atividade': alunos_atividade,
        'modo_edicao': True,
        'motivos_falta': motivos_falta_choices,
    }

    return render(request, 'AppLSD/frequencia_activity_editar.html', context)

def turma_add_alunos(request, turma_id):
    turma = get_object_or_404(Turma, pk=turma_id)

    # base: só alunos frequentando
    
    alunos_queryset = Aluno.objects.filter(
        status_lsd="Frequentando",
        turma__isnull=True,          # <--- só quem não está em nenhuma turma
    )

    # excluir quem já está na turma
    alunos_queryset = alunos_queryset.exclude(turma=turma)

    # turno_oposto: se turma é Matutino, pega alunos Vespertino; se é Vespertino, pega Matutino
    if turma.turno == 'Matutino':
        alunos_queryset = alunos_queryset.filter(turno='Vespertino')
    elif turma.turno == 'Vespertino':
        alunos_queryset = alunos_queryset.filter(turno='Matutino')

    filtro = AlunoFiltroForm(request.GET or None)
    if filtro.is_valid():
        if filtro.cleaned_data['nome']:
            alunos_queryset = alunos_queryset.filter(name__icontains=filtro.cleaned_data['nome'])
        if filtro.cleaned_data['registration_number']:
            alunos_queryset = alunos_queryset.filter(family=filtro.cleaned_data['registration_number'])
        faixa = filtro.cleaned_data.get('faixa_etaria')
        if faixa:
            hoje = date.today()

            def intervalo_idade(min_idade, max_idade):
                # pessoas com idade entre min e max nasceram entre estas datas:
                data_max = date(hoje.year - min_idade, hoje.month, hoje.day)
                data_min = date(hoje.year - max_idade - 1, hoje.month, hoje.day) + timedelta(days=1)
                return data_min, data_max

            if faixa == "06-07":
                data_min, data_max = intervalo_idade(6, 7)
            elif faixa == "08-09":
                data_min, data_max = intervalo_idade(8, 9)
            elif faixa == "10-12":
                data_min, data_max = intervalo_idade(10, 12)
            elif faixa == "13-17":
                data_min, data_max = intervalo_idade(13, 17)

            alunos_queryset = alunos_queryset.filter(
                birth_date__range=(data_min, data_max)
            )
            
    class DinamicoAddAlunosToTurmaForm(AddAlunosToTurmaForm):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self.fields['alunos'].queryset = alunos_queryset

    if request.method == 'POST':
        form = DinamicoAddAlunosToTurmaForm(request.POST)
        if form.is_valid():
            for aluno in form.cleaned_data['alunos']:
                aluno.turma = turma
                aluno.save()
            return redirect('turma_detail', turma_id=turma.id)
    else:
        form = DinamicoAddAlunosToTurmaForm()

    return render(
        request, 
        'AppLSD/turma_add_alunos.html', 
        {'turma': turma, 'form': form, 'filtro': filtro}
    )


@login_required
def mover_aluno(request, aluno_id):
    aluno = get_object_or_404(Aluno, pk=aluno_id)

    # ✅ Proteção real na view (independente do template)
    is_educadora_da_turma = (
        is_educadora(request.user)
        and aluno.turma is not None
        and aluno.turma.educadora == request.user
    )

    if not (request.user.is_superuser or is_coordenacao(request.user) or is_educadora_da_turma):
        raise PermissionDenied

    if request.method == "POST":
        form = MoverAlunoForm(request.POST, user=request.user, aluno=aluno)
        if form.is_valid():
            turma_antiga = aluno.turma
            turma_nova = form.cleaned_data.get('turma_destino')
            motivo = form.cleaned_data.get('motivo')

            aluno.turma = turma_nova
            aluno.save()

            MovimentacaoTurmaAluno.objects.create(
                aluno=aluno,
                turma_origem=turma_antiga,
                turma_destino=turma_nova,
                motivo=motivo
            )

            texto_destino = turma_nova if turma_nova else 'SEM TURMA'
            registrar_log(
                request.user,
                aluno,
                'aluno_movido',
                f'Aluno {aluno.name} movido da turma {turma_antiga} para {texto_destino}. Motivo: {motivo}'
            )

            return redirect('aluno_detail', pk=aluno.pk)
    else:
        form = MoverAlunoForm(initial={'turma_destino': aluno.turma}, user=request.user, aluno=aluno)

    return render(request, "AppLSD/mover_aluno.html", {"form": form, "aluno": aluno})


def adicionar_ocorrencia(request, aluno_id):
    aluno = Aluno.objects.get(pk=aluno_id)
    if request.method == 'POST':
        form = OcorrenciaAlunoForm(request.POST)
        if form.is_valid():
            ocorrencia = form.save(commit=False)
            ocorrencia.aluno = aluno
            ocorrencia.save()
            return redirect('aluno_detail', pk=aluno.pk)
    else:
        form = OcorrenciaAlunoForm()
    return render(request, 'AppLSD/adicionar_ocorrencia.html', {'form': form, 'aluno': aluno})

# HELPERS PARA ADICIONAR ALUNO EM ATIVIDADE ***
def get_turno_atividade_permitido(turno_aluno):
    if turno_aluno == "Matutino":
        return "Vespertino"
    elif turno_aluno == "Vespertino":
        return "Matutino"
    return None


def get_atividades_disponiveis_para_aluno(aluno):
    turno_permitido = get_turno_atividade_permitido(aluno.turno)

    if not turno_permitido:
        return Activity.objects.none()

    return (
        Activity.objects
        .exclude(alunos=aluno)
        .filter(turno=turno_permitido)
        .order_by(Lower("atividade"))
    )

# ADICIONAR ALUNO NA ATIVIDADE EM DETALHES DO ALUNO
@login_required
def aluno_add_atividade(request, aluno_id):
    aluno = get_object_or_404(Aluno, id=aluno_id)

    if request.method != "POST":
        return redirect("aluno_detail", pk=aluno.id)

    if not (request.user.is_superuser or is_coordenacao(request.user) or is_educadora(request.user)):
        messages.error(request, "Sem permissão para adicionar atividade ao aluno.")
        return redirect("aluno_detail", pk=aluno.id)

    atividade_id = request.POST.get("atividade_id")
    if not atividade_id:
        messages.error(request, "Selecione uma atividade.")
        return redirect("aluno_detail", pk=aluno.id)

    atividade = get_object_or_404(Activity, id=atividade_id)

    atividades_disponiveis_ids = list(
        get_atividades_disponiveis_para_aluno(aluno).values_list("id", flat=True)
    )

    if atividade.id not in atividades_disponiveis_ids:
        messages.error(
            request,
            "Esta atividade não é permitida para o turno do aluno ou já está vinculada."
        )
        return redirect("aluno_detail", pk=aluno.id)

    aluno.atividades.add(atividade)

    registrar_log(
        request.user,
        atividade,
        "aluno_adicionado_atividade",
        f'Aluno {aluno.name} adicionado à atividade "{atividade.atividade}".'
    )

    registrar_log(
        request.user,
        aluno,
        "atividade_adicionada_aluno",
        f'Aluno inserido na atividade "{atividade.atividade}".'
    )

    messages.success(request, "Atividade adicionada ao aluno com sucesso.")
    return redirect("aluno_detail", pk=aluno.id)



# ADICIONAR ALUNO NA ATIVIDADE EM DETALHES DA ATIVIDADE ***
@login_required
def activity_add_alunos(request, activity_id):
    atividade = get_object_or_404(Activity, id=activity_id)

    alunos_queryset = (
        Aluno.objects
        .exclude(atividades=atividade)
        .filter(status_lsd="Frequentando")
    )

    if atividade.turno == "Matutino":
        alunos_queryset = alunos_queryset.filter(turno="Vespertino")
    elif atividade.turno == "Vespertino":
        alunos_queryset = alunos_queryset.filter(turno="Matutino")

    filtro = AlunoFiltroForm(request.GET or None)

    if filtro.is_valid():
        nome = filtro.cleaned_data.get("nome")
        inscricao = filtro.cleaned_data.get("registration_number")
        faixa = filtro.cleaned_data.get("faixa_etaria")

        if nome:
            alunos_queryset = alunos_queryset.filter(name__icontains=nome)

        if inscricao:
            alunos_queryset = alunos_queryset.filter(family__registration_number__icontains=inscricao)

        if faixa:
            hoje = date.today()

            def intervalo_idade(min_idade, max_idade):
                data_max = date(hoje.year - min_idade, hoje.month, hoje.day)
                data_min = date(hoje.year - max_idade - 1, hoje.month, hoje.day) + timedelta(days=1)
                return data_min, data_max

            faixas = {
                "06-07": (6, 7),
                "08-09": (8, 9),
                "10-12": (10, 12),
                "13-17": (13, 17),
            }

            if faixa in faixas:
                min_idade, max_idade = faixas[faixa]
                data_min, data_max = intervalo_idade(min_idade, max_idade)
                alunos_queryset = alunos_queryset.filter(
                    birth_date__range=(data_min, data_max)
                )

    alunos_queryset = alunos_queryset.order_by(Lower("name"))

    class DinamicoAddAlunosToAtividadeForm(AddAlunosToTurmaForm):
        def __init__(self, *args, alunos_queryset=None, **kwargs):
            super().__init__(*args, **kwargs)
            self.fields["alunos"].queryset = alunos_queryset or Aluno.objects.none()

    if request.method == "POST":
        form = DinamicoAddAlunosToAtividadeForm(
            request.POST,
            alunos_queryset=alunos_queryset,
        )
        if form.is_valid():
            atividade.alunos.add(*form.cleaned_data["alunos"])
            return redirect("activity_detail", activity_id=atividade.id)
    else:
        form = DinamicoAddAlunosToAtividadeForm(
            alunos_queryset=alunos_queryset,
        )

    context = {
        "atividade": atividade,
        "form": form,
        "filtro": filtro,
    }
    return render(request, "AppLSD/activity_add_alunos.html", context)

def get_contexto_relatorio_turma(request, turma_id, data=None):
    turma = get_object_or_404(Turma, id=turma_id)
    alunos = Aluno.objects.filter(turma=turma).order_by('name')

    data_get = request.GET.get('data')
    if data_get:
        data = data_get

    chamada = None
    presencas_dict = {}
    presentes = 0
    faltantes = 0

    if data:
        if isinstance(data, str):
            try:
                data = datetime.strptime(data, "%Y-%m-%d").date()
            except ValueError:
                data = date.today()

        chamada = FrequenciaTurma.objects.filter(
            turma=turma,
            data=data
        ).first()
    else:
        chamada = (
            FrequenciaTurma.objects
            .filter(turma=turma)
            .order_by('-data')
            .first()
        )
        data = chamada.data if chamada else date.today()

    if chamada:
        presencas = FrequenciaAluno.objects.filter(chamada=chamada).select_related('aluno')
        presencas_dict = {p.aluno_id: p for p in presencas}
        presentes = sum(1 for p in presencas if p.presente)
        faltantes = sum(1 for p in presencas if not p.presente)

    context = {
        'turma': turma,
        'alunos': alunos,
        'chamada': chamada,
        'data': data,
        'presencas_dict': presencas_dict,
        'presentes': presentes,
        'faltantes': faltantes,
        'total_alunos': alunos.count(),
    }
    return context


@login_required
def relatorio_presenca_turma(request, turma_id, data=None):
    context = get_contexto_relatorio_turma(request, turma_id, data)
    return render(request, 'AppLSD/relatorio_presenca_turma.html', context)


@login_required
def relatorio_presenca_turma_pdf(request, turma_id, data=None):
    context = get_contexto_relatorio_turma(request, turma_id, data)

    response = HttpResponse(content_type='application/pdf')
    # response['Content-Disposition'] = f'attachment; filename="relatorio_turma_{turma_id}.pdf"'

    return generate_pdf(
        'AppLSD/relatorio_presenca_turma.html',
        file_object=response,
        context=context,
    )

def get_contexto_relatorio_turma_mensal(request, turma_id):
    turma = Turma.objects.get(id=turma_id)
    alunos = list(
        Aluno.objects
        .filter(turma=turma)
        .order_by('name')
    )

    mes = request.GET.get('mes')
    if not mes:
        hoje = date.today()
        mes = f"{hoje.year}-{hoje.month:02d}"

    ano, mes_num = mes.split('-')
    ano = int(ano)
    mes_num = int(mes_num)

    primeiro_dia = date(ano, mes_num, 1)
    ultimo_dia_num = monthrange(ano, mes_num)[1]
    ultimo_dia = date(ano, mes_num, ultimo_dia_num)

    todos_dias = [
        primeiro_dia + timedelta(days=i)
        for i in range((ultimo_dia - primeiro_dia).days + 1)
    ]

    dias_mes = [d for d in todos_dias if d.weekday() < 5]

    chamadas = (
        FrequenciaTurma.objects
        .filter(turma=turma, data__range=(primeiro_dia, ultimo_dia))
    )
    chamadas_por_data = {c.data: c for c in chamadas}

    frequencias = (
        FrequenciaAluno.objects
        .filter(
            aluno__turma=turma,
            chamada__data__range=(primeiro_dia, ultimo_dia),
        )
        .select_related('aluno', 'chamada')
    )

    freq_dict = {
        (f.aluno_id, f.chamada.data): f
        for f in frequencias
    }

    dias_nao_houve_aula = []
    for dia in dias_mes:
        chamada = chamadas_por_data.get(dia)
        if chamada and chamada.status_aula == FrequenciaTurma.StatusAula.NAO_HOUVE:
            dias_nao_houve_aula.append({
                'data': dia,
                'motivo': chamada.get_motivo_nao_aula_display() if chamada.motivo_nao_aula else '',
                'observacao': chamada.observacao_nao_aula or '',
            })

    linhas = []
    for aluno in alunos:
        linha_status = []
        faltas = 0
        presencas = 0
        justificativas = []

        for dia in dias_mes:
            chamada = chamadas_por_data.get(dia)

            if chamada and chamada.status_aula == FrequenciaTurma.StatusAula.NAO_HOUVE:
                status = 'NH'
            else:
                f = freq_dict.get((aluno.id, dia))

                if f is None:
                    status = ''
                else:
                    if f.presente:
                        status = 'P'
                        presencas += 1
                    else:
                        status = 'F'
                        faltas += 1

                        motivo = (
                            getattr(f, 'get_motivo_falta_display', lambda: '')()
                            or getattr(f, 'motivo_falta', None)
                            or getattr(f, 'justificativa', None)
                            or getattr(f, 'motivo', None)
                            or ''
                        )

                        if motivo:
                            justificativas.append(f"{dia.strftime('%d/%m')}: {motivo}")

            linha_status.append(status)

        total_registros = presencas + faltas
        percentual_presenca = round((presencas / total_registros) * 100, 1) if total_registros else 0

        linhas.append({
            'aluno': aluno,
            'status_por_dia': linha_status,
            'faltas': faltas,
            'presencas': presencas,
            'percentual_presenca': percentual_presenca,
            'justificativas': ' | '.join(justificativas),
        })

    return {
        'turma': turma,
        'mes': mes,
        'ano': ano,
        'mes_num': mes_num,
        'dias_mes': dias_mes,
        'linhas': linhas,
        'dias_nao_houve_aula': dias_nao_houve_aula,
    }

@login_required
def relatorio_turma_mensal_html(request, turma_id):
    context = get_contexto_relatorio_turma_mensal(request, turma_id)
    return render(request, 'AppLSD/relatorio_turma_mensal.html', context)


@login_required
def relatorio_turma_mensal_pdf(request, turma_id):
    context = get_contexto_relatorio_turma_mensal(request, turma_id)
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="relatorio_turma_mensal_{turma_id}.pdf"'
    return generate_pdf(
        'AppLSD/relatorio_turma_mensal_pdf.html',
        file_object=response,
        context=context,
    )


#Relatório diário atividades
def get_contexto_relatorio_atividade(request, activity_id, data=None):
    activity = get_object_or_404(Activity, id=activity_id)

    data_get = request.GET.get('data')
    if data_get:
        data = data_get

    presencas = []
    total_presentes = 0
    total_faltas = 0

    if data:
        presencas = (
            FrequenciaAtividade.objects
            .filter(atividade=activity, data=data)
            .select_related('aluno')
            .order_by('aluno__name')
        )
    else:
        ultima_frequencia = (
            FrequenciaAtividade.objects
            .filter(atividade=activity)
            .order_by('-data')
            .first()
        )

        if ultima_frequencia:
            data = ultima_frequencia.data
            presencas = (
                FrequenciaAtividade.objects
                .filter(atividade=activity, data=data)
                .select_related('aluno')
                .order_by('aluno__name')
            )

    total_presentes = presencas.filter(presente=True).count() if presencas else 0
    total_faltas = presencas.filter(presente=False).count() if presencas else 0

    context = {
        'activity': activity,
        'data': data,
        'presencas': presencas,
        'total_presentes': total_presentes,
        'total_faltas': total_faltas,
    }
    return context


@login_required
def relatorio_presenca_activity(request, activity_id, data=None):
    context = get_contexto_relatorio_atividade(request, activity_id, data)
    return render(request, 'AppLSD/relatorio_presenca_activity.html', context)


@login_required
def relatorio_presenca_activity_pdf(request, activity_id, data=None):
    context = get_contexto_relatorio_atividade(request, activity_id, data)
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="relatorio_atividade_{activity_id}.pdf"'
    return generate_pdf(
        'AppLSD/relatorio_presenca_activity.html',
        file_object=response,
        context=context,
    )

#Relatório Mensal atividdes
def normalizar_texto(texto):
    if not texto:
        return ''
    texto = str(texto).strip().lower()
    texto = unicodedata.normalize('NFKD', texto)
    texto = ''.join(c for c in texto if not unicodedata.combining(c))
    return texto

def extrair_weekdays(dia_semana_raw):
    valor = normalizar_texto(dia_semana_raw)

    for sep in [';', '|', '/', ' e ']:
        valor = valor.replace(sep, ',')

    partes = [p.strip() for p in valor.split(',') if p.strip()]

    mapa_dias = {
        'segunda': 0,
        'segunda-feira': 0,
        'seg': 0,
        'segunda feira': 0,

        'terca': 1,
        'terca-feira': 1,
        'ter': 1,
        'terca feira': 1,

        'quarta': 2,
        'quarta-feira': 2,
        'qua': 2,
        'quarta feira': 2,

        'quinta': 3,
        'quinta-feira': 3,
        'qui': 3,
        'quinta feira': 3,

        'sexta': 4,
        'sexta-feira': 4,
        'sex': 4,
        'sexta feira': 4,

        'sabado': 5,
        'sab': 5,
        'sabado-feira': 5,
        'sabado feira': 5,

        'domingo': 6,
        'dom': 6,
    }

    weekdays = set()

    for parte in partes:
        parte = re.sub(r'\s+', ' ', parte).strip()

        if parte in mapa_dias:
            weekdays.add(mapa_dias[parte])
            continue

        for chave, numero in mapa_dias.items():
            if chave in parte:
                weekdays.add(numero)
                break

    return weekdays


def get_contexto_relatorio_atividade_mensal(request, activity_id):
    activity = get_object_or_404(Activity, id=activity_id)

    mes = request.GET.get('mes')
    if not mes:
        hoje = date.today()
        mes = f"{hoje.year}-{hoje.month:02d}"

    ano, mes_num = mes.split('-')
    ano = int(ano)
    mes_num = int(mes_num)

    primeiro_dia = date(ano, mes_num, 1)
    ultimo_dia_num = monthrange(ano, mes_num)[1]
    ultimo_dia = date(ano, mes_num, ultimo_dia_num)

    todos_dias = [
        primeiro_dia + timedelta(days=i)
        for i in range((ultimo_dia - primeiro_dia).days + 1)
    ]

    weekdays_permitidos = extrair_weekdays(activity.dia_semana)

    dias_mes = [
        d for d in todos_dias
        if d.weekday() in weekdays_permitidos
    ]

    if not dias_mes:
        dias_mes = [d for d in todos_dias if d.weekday() < 5]

    alunos = list(activity.alunos.all().order_by('name'))

    frequencias = (
        FrequenciaAtividade.objects
        .filter(
            atividade=activity,
            data__range=(primeiro_dia, ultimo_dia),
            aluno__in=alunos,
        )
        .select_related('aluno')
        .order_by('data', 'aluno__name')
    )

    freq_dict = {
        (f.aluno_id, f.data): f
        for f in frequencias
    }

    linhas = []
    for aluno in alunos:
        linha_status = []
        faltas = 0
        presencas = 0
        justificativas = []

        for dia in dias_mes:
            f = freq_dict.get((aluno.id, dia))

            if f is None:
                status = ''
            else:
                if f.presente:
                    status = 'P'
                    presencas += 1
                else:
                    status = 'F'
                    faltas += 1
                    if f.motivo_falta:
                        justificativas.append(f"{dia.strftime('%d/%m')}: {f.motivo_falta}")

            linha_status.append(status)

        total_registros = presencas + faltas
        percentual_presenca = round((presencas / total_registros) * 100, 1) if total_registros else 0

        linhas.append({
            'aluno': aluno,
            'status_por_dia': linha_status,
            'faltas': faltas,
            'presencas': presencas,
            'percentual_presenca': percentual_presenca,
            'justificativas': ' | '.join(justificativas),
        })

    return {
        'activity': activity,
        'mes': mes,
        'ano': ano,
        'mes_num': mes_num,
        'dias_mes': dias_mes,
        'linhas': linhas,
    }

@login_required
def relatorio_atividade_mensal_html(request, activity_id):
    context = get_contexto_relatorio_atividade_mensal(request, activity_id)
    return render(request, 'AppLSD/relatorio_atividade_mensal.html', context)


@login_required
def relatorio_atividade_mensal_pdf(request, activity_id):
    context = get_contexto_relatorio_atividade_mensal(request, activity_id)
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="relatorio_atividade_mensal_{activity_id}.pdf"'
    return generate_pdf(
        'AppLSD/relatorio_atividade_mensal_pdf.html',
        file_object=response,
        context=context,
    )



def get_contexto_relatorio_aluno_mensal(request, aluno_id=None):
    aluno_id = aluno_id or request.GET.get('aluno_id')
    mes = request.GET.get('mes')  # YYYY-MM

    if not aluno_id:
        raise Http404("Aluno não informado.")

    aluno = get_object_or_404(Aluno, id=aluno_id)

    if not mes:
        hoje = date.today()
        mes = f"{hoje.year}-{hoje.month:02d}"

    try:
        ano, mes_num = mes.split('-')
        ano = int(ano)
        mes_num = int(mes_num)
        primeiro_dia = date(ano, mes_num, 1)
    except (ValueError, TypeError):
        hoje = date.today()
        ano = hoje.year
        mes_num = hoje.month
        mes = f"{ano}-{mes_num:02d}"
        primeiro_dia = date(ano, mes_num, 1)

    ultimo_dia_num = monthrange(ano, mes_num)[1]
    ultimo_dia = date(ano, mes_num, ultimo_dia_num)

    if mes_num == 1:
        mes_anterior = f"{ano - 1}-12"
    else:
        mes_anterior = f"{ano}-{mes_num - 1:02d}"

    if mes_num == 12:
        mes_proximo = f"{ano + 1}-01"
    else:
        mes_proximo = f"{ano}-{mes_num + 1:02d}"

    todos_dias = [
        primeiro_dia + timedelta(days=i)
        for i in range((ultimo_dia - primeiro_dia).days + 1)
    ]

    dias_mes = [d for d in todos_dias if d.weekday() < 5]

    turma_aluno = getattr(aluno, 'turma', None)

    freq_qs = (
        FrequenciaAluno.objects
        .filter(
            aluno=aluno,
            chamada__data__range=(primeiro_dia, ultimo_dia),
        )
        .select_related('chamada')
        .order_by('chamada__data')
    )

    freq_dict = {f.chamada.data: f for f in freq_qs}

    status_por_dia = []
    faltas = 0
    presentes = 0

    for dia in dias_mes:
        f = freq_dict.get(dia)

        if f is None:
            status = ''
        else:
            if f.presente:
                status = 'P'
                presentes += 1
            else:
                status = 'F'
                faltas += 1

        status_por_dia.append({
            'dia': dia,
            'status': status,
        })

    atividades_freq_qs = (
        FrequenciaAtividade.objects
        .filter(
            aluno=aluno,
            data__range=(primeiro_dia, ultimo_dia),
        )
        .select_related('atividade', 'atividade__facilitador')
        .order_by('atividade__atividade', 'data')
    )

    freq_atividade_dict = {}
    for reg in atividades_freq_qs:
        freq_atividade_dict[(reg.atividade_id, reg.data)] = reg

    atividades_grade = []

    atividades_vinculadas = (
        aluno.atividades.all()
        .select_related('facilitador')
        .order_by('atividade')
    )

    for atividade in atividades_vinculadas:
        dias_atividade = []
        faltas_atividade = 0

        for dia in dias_mes:
            incluir_dia = True

            if getattr(atividade, 'dia_semana', None):
                mapa_dias = {
                    'segunda': 0,
                    'terça': 1,
                    'terca': 1,
                    'quarta': 2,
                    'quinta': 3,
                    'sexta': 4,
                    'sábado': 5,
                    'sabado': 5,
                    'domingo': 6,
                }

                dia_semana_atividade = str(atividade.dia_semana).lower()
                incluir_dia = False

                for nome, numero in mapa_dias.items():
                    if nome in dia_semana_atividade and dia.weekday() == numero:
                        incluir_dia = True
                        break

            if incluir_dia:
                reg = freq_atividade_dict.get((atividade.id, dia))
                if reg:
                    if reg.presente:
                        status = 'P'
                    else:
                        status = 'F'
                        faltas_atividade += 1
                else:
                    status = ''
            else:
                status = None

            dias_atividade.append({
                'dia': dia,
                'status': status,
            })

        atividades_grade.append({
            'atividade': atividade,
            'dias': dias_atividade,
            'faltas': faltas_atividade,
        })

    movs = (
        MovimentacaoTurmaAluno.objects
        .filter(
            aluno=aluno,
            data__date__range=(primeiro_dia, ultimo_dia),
        )
        .select_related('turma_origem', 'turma_destino')
        .order_by('data')
    )

    ocorrencias = (
        OcorrenciaAluno.objects
        .filter(
            aluno=aluno,
            data__range=(primeiro_dia, ultimo_dia),
        )
        .order_by('data')
    )

    context = {
        'aluno': aluno,
        'mes': mes,
        'ano': ano,
        'mes_num': mes_num,
        'mes_anterior': mes_anterior,
        'mes_proximo': mes_proximo,
        'dias_mes': dias_mes,
        'status_por_dia': status_por_dia,
        'faltas': faltas,
        'presentes': presentes,
        'total_dias_letivos': len(dias_mes),
        'atividades_grade': atividades_grade,
        'movimentacoes': movs,
        'ocorrencias': ocorrencias,
        'turma_relatorio': turma_aluno,
        'aluno_sem_turma': turma_aluno is None,
    }
    return context


@login_required
def relatorio_mensal_aluno(request):
    context = get_contexto_relatorio_aluno_mensal(request)
    return render(request, 'AppLSD/relatorio_mensal_aluno.html', context)


#@login_required
def relatorio_busca_aluno(request):
    alunos = Aluno.objects.order_by('name')
    return render(request, 'AppLSD/relatorio_busca_aluno.html', {'alunos': alunos})

def relatorio_grafico_mensal_aluno(request):
    aluno_id = request.GET.get('aluno_id')
    mes_ano = request.GET.get('mes')  # formato "YYYY-MM"
    aluno = get_object_or_404(Aluno, pk=aluno_id)

    ano, mes = (int(x) for x in mes_ano.split('-'))
    chamadas = FrequenciaTurma.objects.filter(aluno=aluno, data__year=ano, data__month=mes)

    presencas = sum(1 for c in chamadas if c.presente)
    faltas = sum(1 for c in chamadas if not c.presente)

    # Criar gráfico
    labels = ['Presenças', 'Faltas']
    values = [presencas, faltas]
    fig, ax = plt.subplots()
    ax.bar(labels, values, color=['green', 'red'])
    ax.set_title(f'Frequência de {aluno.name} - {mes_ano}')

    # Convertendo gráfico para imagem no formato base64
    buf = io.BytesIO()
    plt.savefig(buf, format='png')
    buf.seek(0)
    image_png = buf.getvalue()
    graph = base64.b64encode(image_png).decode('utf-8')
    buf.close()

    return render(request, 'AppLSD/relatorio_grafico_aluno.html', {
        'graph': graph,
        'aluno': aluno,
        'mes_ano': mes_ano,
        'presencas': presencas,
        'faltas': faltas,
    })

@login_required
def exportar_frequencia_mensal_aluno_excel(request):
    aluno_id = request.GET.get('aluno_id')
    mes_ano = request.GET.get('mes')
    aluno = get_object_or_404(Aluno, pk=aluno_id)
    ano, mes = (int(x) for x in mes_ano.split('-'))
    chamadas = FrequenciaTurma.objects.filter(aluno=aluno, data__year=ano, data__month=mes)

    wb = Workbook()
    ws = wb.active
    ws.title = "Frequência"

    ws.append(["Data", "Presente", "Observação"])

    for c in chamadas:
        ws.append([c.data.strftime('%d/%m/%Y'), "Sim" if c.presente else "Não", c.observacao or ""])

    response = HttpResponse(content_type="application/ms-excel")
    response["Content-Disposition"] = f"attachment; filename=frequencia_{aluno.name}_{mes_ano}.xlsx"
    wb.save(response)
    return response

""" 
@login_required
def home_relatorios(request):
    turmas = Turma.objects.all()
    atividades = Activity.objects.all()
    alunos = Aluno.objects.all()  # Incluído para o relatório de aluno
    return render(request, 'AppLSD/home_relatorios.html', {
        'turmas': turmas,
        'atividades': atividades,
        'alunos': alunos,
    })
"""

def registrar_faltas_automaticas(chamada):
    """
    Registra automaticamente como falta todos os alunos da turma
    que não foram marcados como presentes na chamada.
    
    Args:
        chamada: Objeto FrequenciaTurma
    """
        
    # Buscar todos os alunos da turma
    alunos_turma = Aluno.objects.filter(turma=chamada.turma)
    
    # Alunos que já têm registro nesta chamada
    alunos_com_registro = FrequenciaAluno.objects.filter(
        chamada=chamada
    ).values_list('aluno_id', flat=True)
    
    # Alunos sem registro = faltaram
    alunos_sem_registro = alunos_turma.exclude(id__in=alunos_com_registro)
    
    # Registrar as faltas
    faltas_criadas = 0
    for aluno in alunos_sem_registro:
        FrequenciaAluno.objects.create(
            chamada=chamada,
            aluno=aluno,
            presente=False
        )
        faltas_criadas += 1
    
    return faltas_criadas

def finalizar_chamada(request, chamada_id):
    chamada = FrequenciaTurma.objects.get(id=chamada_id)
    
    # Registrar faltas automaticamente
    faltas_criadas = registrar_faltas_automaticas(chamada)
    
    messages.success(request, f"Chamada finalizada! {faltas_criadas} faltas registradas automaticamente.")
    return redirect(request, 'AppLSD/home_relatorios.html')

def get_created_at_for_instance(instance, acao):
    """
    Retorna a data/hora do primeiro log da ação informada
    para a instância (por ex. 'familia_criada').
    """
    ct = ContentType.objects.get_for_model(instance.__class__)
    log = (
        AppLog.objects
        .filter(
            content_type=ct,
            object_id=instance.pk,
            acao=acao,
        )
        .order_by('criado_em')
        .first()
    )
    return log.criado_em if log else None

#Exportação em excel das listas Família, Aluno e Adulto
@login_required
def export_family_excel(request):
    campos = request.GET.getlist('campos')

    if not campos:
        campos = [
            'criado_em', 'registration_number', 'responsible_name', 'cpf',
            'telephone', 'address', 'neighborhood', 'status'
        ]

    families = Family.objects.all().order_by('registration_number')

    wb = Workbook()
    ws = wb.active
    ws.title = "Famílias"

    campos_info = {
        'criado_em': {'label': 'Criado em'},  # novo
        'registration_number': {'label': 'Inscrição', 'field': 'registration_number'},
        'responsible_name': {'label': 'Responsável', 'field': 'responsible_name'},
        'cpf': {'label': 'CPF', 'field': 'cpf'},
        'rg': {'label': 'RG', 'field': 'rg'},
        'nis': {'label': 'NIS', 'field': 'nis'},
        'birth_date': {'label': 'Data de Nascimento', 'field': 'birth_date'},
        'sex': {'label': 'Sexo', 'field': 'sex'},
        'telephone': {'label': 'Telefone', 'field': 'telephone'},
        'telephone_2': {'label': 'Telefone 2', 'field': 'telephone_2'},
        'address': {'label': 'Endereço', 'field': 'address'},
        'number': {'label': 'Número', 'field': 'number'},
        'neighborhood': {'label': 'Bairro', 'field': 'neighborhood'},
        'cep': {'label': 'CEP', 'field': 'cep'},
        'reference_point': {'label': 'Ponto de Referência', 'field': 'reference_point'},
        'marital_status': {'label': 'Estado Civil', 'field': 'marital_status'},
        'education': {'label': 'Escolaridade', 'field': 'education'},
        'race': {'label': 'Raça', 'field': 'race'},
        'occupation': {'label': 'Ocupação', 'field': 'occupation'},
        'salary_range': {'label': 'Faixa Salarial', 'field': 'salary_range'},
        'num_residents': {'label': 'Nº de Moradores', 'field': 'num_residents'},
        'status': {'label': 'Status', 'field': 'status'},
    }

    headers = [campos_info[c]['label'] for c in campos if c in campos_info]
    ws.append(headers)

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    for family in families:
        row = []
        for campo in campos:
            if campo not in campos_info:
                continue

            valor = ''
            if campo == 'criado_em':
                dt = get_created_at_for_instance(family, 'familia_criada')
                valor = dt.strftime('%d/%m/%Y %H:%M') if dt else ''
            else:
                field_name = campos_info[campo]['field']
                v = getattr(family, field_name, '')
                if campo == 'birth_date' and v:
                    v = v.strftime('%d/%m/%Y')
                valor = v

            row.append(str(valor) if valor else '')

        ws.append(row)

    # (restante da função igual: ajustar colunas, salvar response)
# Ajusta largura das colunas
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(cell.value)
            except:
                pass
        adjusted_width = min((max_length + 2), 50)
        ws.column_dimensions[column_letter].width = adjusted_width
   
    # Resposta HTTP
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f'familias_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
   
    wb.save(response)
    return response


@login_required
def export_aluno_excel(request):
    campos = request.GET.getlist('campos')
    
       
    alunos = Aluno.objects.select_related('family', 'turma').all().order_by('name')
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Alunos"
    
    # Dicionário com TODOS os campos possíveis
    campos_info = {
        'criado_em': {'label': 'Criado em'},  # novo
        'inscricao': {'label': 'Inscrição', 'field': 'family__registration_number'},
        'name': {'label': 'Nome', 'field': 'name'},
        'responsavel': {'label': 'Responsável', 'field': 'family__responsible_name'},
        'cpf': {'label': 'CPF', 'field': 'cpf'},
        'nis': {'label': 'NIS', 'field': 'nis'},
        'birth_date': {'label': 'Data de Nascimento', 'field': 'birth_date'},
        'idade': {'label': 'Idade', 'field': 'idade'},
        'sex': {'label': 'Sexo', 'field': 'sex'},
        'rede_ensino': {'label': 'Rede de Ensino', 'field': 'rede_ensino'},
        'school': {'label': 'Escola', 'field': 'school'},
        'ensino': {'label': 'Ensino', 'field': 'ensino'},
        'serie': {'label': 'Série', 'field': 'serie'},
        'turno': {'label': 'Turno', 'field': 'turno'},
        'health_problem': {'label': 'Tem Problema de Saúde?', 'field': 'health_problem'},
        'special_need': {'label': 'Qual problema de Saúde?', 'field': 'special_need'},
        'uso_medicacao': {'label': 'Faz uso de Medicação?', 'field': 'uso_medicacao'},
        'qual_medicacao': {'label': 'Qual Medicação?', 'field': 'qual_medicacao'},
        'frequencia_tipo': {'label': 'Tipo de Frequência', 'field': 'frequencia_tipo'},
        'dias_semana': {'label': 'Dias da Semana', 'field': 'dias_semana'},
        'status_lsd': {'label': 'Status LSD', 'field': 'status_lsd'},
        'turma': {'label': 'Turma', 'field': 'turma__name'},
    }

    if not campos:
        campos = list(campos_info.keys())  # se não vier campos, exporta todos os campos disponíveis
    
    # Cabeçalhos NA ORDEM dos campos selecionados
    headers = []
    for campo in campos:
        if campo in campos_info:
            headers.append(campos_info[campo]['label'])
    
    ws.append(headers)
    
    # Estilo do cabeçalho
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")
    
    # Dados NA ORDEM dos campos selecionados
    for aluno in alunos:
        row = []
        for campo in campos:
            if campo in campos_info:
                valor = ''
                
                # Campos especiais
                if campo == 'inscricao':
                    valor = aluno.family.registration_number if aluno.family else ''
                elif campo == 'responsavel':
                    valor = aluno.family.responsible_name if aluno.family else ''
                elif campo == 'birth_date':
                    valor = aluno.birth_date.strftime('%d/%m/%Y') if aluno.birth_date else ''
                elif campo == 'idade':
                    valor = aluno.idade if hasattr(aluno, 'idade') else ''
                elif campo == 'criado_em':
                    dt = get_created_at_for_instance(aluno, 'aluno_criado')
                    valor = dt.strftime('%d/%m/%Y %H:%M') if dt else ''
                else:
                    valor = getattr(aluno, campo, '')
                
                row.append(str(valor) if valor else '')
        
        ws.append(row)
    
    # Ajusta largura das colunas
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(cell.value)
            except:
                pass
        adjusted_width = min((max_length + 2), 50)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f'alunos_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    
    wb.save(response)
    return response

@login_required
def export_adult_excel(request):
     # DEBUG - Lista todos os campos do modelo
    adult = Adult.objects.first()
    if adult:
        print("===== CAMPOS DISPONÍVEIS NO MODELO ADULT =====")
        for field in adult._meta.get_fields():
            print(f"Campo: {field.name}")
        print("=" * 50)
    campos = request.GET.getlist('campos')
    
    if not campos:
        campos = ['criado_em', 'inscricao', 'name', 'cpf', 'birth_date', 'parentesco', 'ocupacao']
    
    adults = Adult.objects.select_related('family').all().order_by('name')
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Adultos"
    
    # Dicionário com TODOS os campos possíveis
    campos_info = {
        'criado_em': {'label': 'Criado em'},
        'inscricao': {'label': 'Inscrição', 'field': 'family__registration_number'},
        'name': {'label': 'Nome', 'field': 'name'},
        'cpf': {'label': 'CPF', 'field': 'cpf'},
        'birth_date': {'label': 'Data de Nascimento', 'field': 'birth_date'},
        'idade': {'label': 'Idade', 'field': 'idade'},
        'sex': {'label': 'Sexo', 'field': 'sex'},
        'education': {'label': 'Escolaridade', 'field': 'education'},
        'parentesco': {'label': 'Parentesco', 'field': 'parentesco'},
        'ocupacao': {'label': 'Ocupação', 'field': 'ocupacao'},
        'renda': {'label': 'Renda', 'field': 'income'},
        'telephone': {'label': 'Telefone', 'field': 'telephone'},
        'responsavel': {'label': 'Responsável', 'field': 'family__responsible_name'},
    }
    
    # Cabeçalhos NA ORDEM dos campos selecionados
    headers = []
    for campo in campos:
        if campo in campos_info:
            headers.append(campos_info[campo]['label'])
    
    ws.append(headers)
    
    # Estilo do cabeçalho
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")
    
    # Dados NA ORDEM dos campos selecionados
    for adult in adults:
        row = []
        for campo in campos:
            if campo in campos_info:
                valor = ''
                
                # Campos especiais
                if campo == 'inscricao':
                    valor = adult.family.registration_number if adult.family else ''
                elif campo == 'responsavel':
                    valor = adult.family.responsible_name if adult.family else ''
                elif campo == 'birth_date':
                    valor = adult.birth_date.strftime('%d/%m/%Y') if adult.birth_date else ''
                elif campo == 'idade':
                    valor = adult.idade if hasattr(adult, 'idade') else ''
                elif campo == 'is_working':
                    valor = 'Sim' if adult.is_working == 'sim' else 'Não' if adult.is_working else ''
                elif campo == 'criado_em':
                    dt = get_created_at_for_instance(adult, 'adulto_criado')
                    valor = dt.strftime('%d/%m/%Y %H:%M') if dt else ''
                else:
                    valor = getattr(adult, campo, '')
                
                row.append(str(valor) if valor else '')
        
        ws.append(row)
    
    # Ajusta largura das colunas
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(cell.value)
            except:
                pass
        adjusted_width = min((max_length + 2), 50)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f'adultos_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    
    wb.save(response)
    return response



def _get_periodo(request):
    hoje = date.today()
    mes_param = request.GET.get("mes")
    if mes_param:
        ano, mes = map(int, mes_param.split("-"))
    else:
        ano, mes = hoje.year, hoje.month

    primeiro_dia = date(ano, mes, 1)
    ultimo_dia_num = monthrange(ano, mes)[1]
    ultimo_dia = date(ano, mes, ultimo_dia_num)
    return ano, mes, primeiro_dia, ultimo_dia


def _auto_width(ws, max_width=40):
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
        ws.column_dimensions[col_letter].width = min(max_len + 2, max_width)


@login_required
def relatorio_busca_ativa(request):
    ano, mes, primeiro_dia, ultimo_dia = _get_periodo(request)

    alunos = (
        Aluno.objects
        .filter(
            status_lsd__iexact="Frequentando",
            excluido=False,
            family__excluido=False,
        )
        .select_related("family", "turma", "turma__educadora")
        .annotate(
            familia=F("family__responsible_name"),
            contato=F("family__telephone"),
            endereco=Concat(
                Coalesce(F("family__address"), Value("")),
                Value(", "),
                Coalesce(F("family__number"), Value("")),
                Value(", "),
                Coalesce(F("family__neighborhood"), Value("")),
                Value(", "),
                Coalesce(F("family__reference_point"), Value("")),
                output_field=CharField()
            ),
            total_chamadas=Coalesce(
                Sum(
                    Case(
                        When(
                            frequenciaaluno__chamada__data__range=(primeiro_dia, ultimo_dia),
                            then=1
                        ),
                        default=0,
                        output_field=IntegerField()
                    )
                ),
                0
            ),
            total_presencas=Coalesce(
                Sum(
                    Case(
                        When(
                            frequenciaaluno__chamada__data__range=(primeiro_dia, ultimo_dia),
                            frequenciaaluno__presente=True,
                            then=1
                        ),
                        default=0,
                        output_field=IntegerField()
                    )
                ),
                0
            ),
        )
        .annotate(
            percentual_presenca=Case(
                When(
                    total_chamadas__gt=0,
                    then=(F("total_presencas") * 100.0 / F("total_chamadas"))
                ),
                default=Value(0.0),
                output_field=FloatField()
            )
        )
        .filter(percentual_presenca__lt=50)
        .order_by("family__responsible_name", "name")
    )

    familias = defaultdict(list)
    for aluno in alunos:
        chave = (aluno.familia, aluno.contato, aluno.endereco)
        familias[chave].append(aluno)

    wb = Workbook()
    ws = wb.active
    ws.title = "Busca Ativa"

    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    cab_fill = PatternFill("solid", fgColor="D9D9D9")
    sub_fill = PatternFill("solid", fgColor="EDEDED")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 40

    row = 1
    for (responsavel, contato, endereco), assistidos in familias.items():
        ws.cell(row=row, column=1, value="Responsável")
        ws.cell(row=row, column=2, value="Contato")
        ws.cell(row=row, column=3, value="Endereço")
        for c in range(1, 4):
            ws.cell(row=row, column=c).font = Font(bold=True)
            ws.cell(row=row, column=c).fill = cab_fill
            ws.cell(row=row, column=c).alignment = center
            ws.cell(row=row, column=c).border = border

        row += 1
        ws.cell(row=row, column=1, value=responsavel)
        ws.cell(row=row, column=2, value=contato)
        ws.cell(row=row, column=3, value=endereco)
        for c in range(1, 4):
            ws.cell(row=row, column=c).alignment = left
            ws.cell(row=row, column=c).border = border

        row += 1
        ws.cell(row=row, column=1, value="Assistidos")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        ws.cell(row=row, column=3, value="Percentual presença")

        for c in range(1, 4):
            ws.cell(row=row, column=c).font = Font(bold=True)
            ws.cell(row=row, column=c).fill = sub_fill
            ws.cell(row=row, column=c).alignment = center
            ws.cell(row=row, column=c).border = border

        row += 1
        for assistido in assistidos:
            ws.cell(row=row, column=1, value=assistido.name)
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
            ws.cell(row=row, column=3, value=round(assistido.percentual_presenca, 1))

            for c in range(1, 4):
                ws.cell(row=row, column=c).border = border
                ws.cell(row=row, column=c).alignment = left if c != 3 else center

            row += 1

        row += 1

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="busca_ativa_{ano}_{mes:02d}.xlsx"'
    return response


@login_required
def relatorio_busca_ativa_html(request):
    ano, mes, primeiro_dia, ultimo_dia = _get_periodo(request)

    alunos = (
        Aluno.objects
        .filter(
            status_lsd__iexact="Frequentando",
            excluido=False,
            family__excluido=False,
        )
        .select_related("family")
        .annotate(
            familia=F("family__responsible_name"),
            contato=F("family__telephone"),
            endereco=Concat(
                Coalesce(F("family__address"), Value("")),
                Value(", "),
                Coalesce(F("family__number"), Value("")),
                Value(", "),
                Coalesce(F("family__neighborhood"), Value("")),
                Value(", "),
                Coalesce(F("family__reference_point"), Value("")),
                output_field=CharField()
            ),
            total_chamadas=Coalesce(
                Sum(
                    Case(
                        When(
                            frequenciaaluno__chamada__data__range=(primeiro_dia, ultimo_dia),
                            then=1
                        ),
                        default=0,
                        output_field=IntegerField()
                    )
                ),
                0
            ),
            total_presencas=Coalesce(
                Sum(
                    Case(
                        When(
                            frequenciaaluno__chamada__data__range=(primeiro_dia, ultimo_dia),
                            frequenciaaluno__presente=True,
                            then=1
                        ),
                        default=0,
                        output_field=IntegerField()
                    )
                ),
                0
            ),
        )
        .annotate(
            percentual_presenca=Case(
                When(total_chamadas__gt=0, then=(F("total_presencas") * 100.0 / F("total_chamadas"))),
                default=Value(0.0),
                output_field=FloatField()
            )
        )
        .filter(percentual_presenca__lt=50)
        .order_by("family__responsible_name", "name")
    )

    familias = defaultdict(list)
    for aluno in alunos:
        chave = {
            "responsavel": aluno.familia,
            "contato": aluno.contato,
            "endereco": aluno.endereco,
        }
        familias[(aluno.familia, aluno.contato, aluno.endereco)].append({
            "assistido": aluno.name,
            "percentual_presenca": round(aluno.percentual_presenca, 1)
        })

    blocos = []
    for (responsavel, contato, endereco), assistidos in familias.items():
        blocos.append({
            "responsavel": responsavel,
            "contato": contato,
            "endereco": endereco,
            "assistidos": assistidos,
        })

    context = {
        "titulo": "Relatório Busca Ativa",
        "mes": f"{ano}-{mes:02d}",
        "blocos": blocos,
    }
    return render(request, "AppLSD/relatorio_busca_ativa.html", context)


@login_required
def relatorio_sisc(request):
    ano, mes, primeiro_dia, ultimo_dia = _get_periodo(request)

    todos_dias = [
        primeiro_dia + timedelta(days=i)
        for i in range((ultimo_dia - primeiro_dia).days + 1)
    ]
    dias_mes = [d for d in todos_dias if d.weekday() < 5]

    alunos = list(
        Aluno.objects
        .filter(
            status_lsd__iexact="Frequentando",
            excluido=False,
            family__excluido=False,
            turma__isnull=False,
            turma__educadora__isnull=False,
        )
        .select_related("family", "turma", "turma__educadora")
        .order_by(
            "turma__educadora__first_name",
            "turma__educadora__last_name",
            "turma__turno",
            "name"
        )
    )

    frequencias = (
        FrequenciaAluno.objects
        .filter(
            aluno__in=alunos,
            chamada__data__range=(primeiro_dia, ultimo_dia),
        )
        .select_related("aluno", "chamada")
    )

    freq_dict = {(f.aluno_id, f.chamada.data): f for f in frequencias}

    wb = Workbook()
    ws = wb.active
    ws.title = "Relatorio SISC"

    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    cab_fill = PatternFill("solid", fgColor="0F766E")
    cab_font = Font(color="FFFFFF", bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    cabecalho = ["Educadora", "Turma", "Aluno"]
    cabecalho += [d.strftime("%d") for d in dias_mes]
    cabecalho += ["Total Faltas", "% Presença", "Justificativas"]

    ws.append(cabecalho)

    for c in range(1, len(cabecalho) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = cab_fill
        cell.font = cab_font
        cell.alignment = center
        cell.border = border

    row = 2
    for aluno in alunos:
        educadora = f"{aluno.turma.educadora.first_name} {aluno.turma.educadora.last_name}".strip()
        turma_nome = " - ".join(
            [v for v in [aluno.turma.grupo, aluno.turma.turno, aluno.turma.faixa_etaria] if v]
        )

        linha = [educadora, turma_nome, aluno.name]
        total_presencas = 0
        total_faltas = 0
        justificativas = []

        for dia in dias_mes:
            f = freq_dict.get((aluno.id, dia))
            if f is None:
                status = ""
            elif f.presente:
                status = "P"
                total_presencas += 1
            else:
                status = "F"
                total_faltas += 1
                motivo = getattr(f, "motivo_falta", "") or getattr(f, "justificativa", "") or ""
                if motivo:
                    justificativas.append(f"{dia.strftime('%d/%m')}: {motivo}")

            linha.append(status)

        total_chamadas = total_presencas + total_faltas
        percentual_presenca = round((total_presencas / total_chamadas) * 100, 1) if total_chamadas else 0

        linha.append(total_faltas)
        linha.append(percentual_presenca)
        linha.append(" | ".join(justificativas))

        ws.append(linha)

        for c in range(1, len(cabecalho) + 1):
            ws.cell(row=row, column=c).border = border
            ws.cell(row=row, column=c).alignment = center if 4 <= c <= (3 + len(dias_mes)) else left

        row += 1

    ws.freeze_panes = "D2"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 32
    ws.column_dimensions[get_column_letter(len(cabecalho) - 2)].width = 12
    ws.column_dimensions[get_column_letter(len(cabecalho) - 1)].width = 12
    ws.column_dimensions[get_column_letter(len(cabecalho))].width = 40

    for i in range(4, 4 + len(dias_mes)):
        ws.column_dimensions[get_column_letter(i)].width = 4

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="relatorio_sisc_{ano}_{mes:02d}.xlsx"'
    return response


@login_required
def relatorio_sisc_html(request):
    ano, mes, primeiro_dia, ultimo_dia = _get_periodo(request)

    todos_dias = [
        primeiro_dia + timedelta(days=i)
        for i in range((ultimo_dia - primeiro_dia).days + 1)
    ]
    dias_mes = [d for d in todos_dias if d.weekday() < 5]

    alunos = list(
        Aluno.objects
        .filter(
            status_lsd__iexact="Frequentando",
            excluido=False,
            family__excluido=False,
            turma__isnull=False,
            turma__educadora__isnull=False,
        )
        .select_related("turma", "turma__educadora")
        .order_by(
            "turma__educadora__first_name",
            "turma__educadora__last_name",
            "turma__turno",
            "name"
        )
    )

    frequencias = (
        FrequenciaAluno.objects
        .filter(
            aluno__in=alunos,
            chamada__data__range=(primeiro_dia, ultimo_dia),
        )
        .select_related("aluno", "chamada")
    )

    freq_dict = {(f.aluno_id, f.chamada.data): f for f in frequencias}

    linhas = []
    educadoras = set()

    for aluno in alunos:
        educadora = f"{aluno.turma.educadora.first_name} {aluno.turma.educadora.last_name}".strip()
        educadoras.add(educadora)

        status_por_dia = []
        faltas = 0
        presencas = 0
        justificativas = []

        for dia in dias_mes:
            f = freq_dict.get((aluno.id, dia))
            if f is None:
                status = ""
            elif f.presente:
                status = "P"
                presencas += 1
            else:
                status = "F"
                faltas += 1
                motivo = getattr(f, "motivo_falta", "") or getattr(f, "justificativa", "") or ""
                if motivo:
                    justificativas.append(f"{dia.strftime('%d/%m')}: {motivo}")
            status_por_dia.append(status)

        total = presencas + faltas
        percentual = round((presencas / total) * 100, 1) if total else 0

        linhas.append({
            "educadora": educadora,
            "aluno": aluno.name,
            "status_por_dia": status_por_dia,
            "faltas": faltas,
            "percentual": percentual,
            "justificativas": " | ".join(justificativas),
        })

    context = {
        "titulo": "Relatório SISC",
        "mes": f"{ano}-{mes:02d}",
        "educadora_texto": " / ".join(sorted(educadoras)),
        "dias_mes": dias_mes,
        "linhas": linhas,
    }
    return render(request, "AppLSD/relatorio_sisc.html", context)


def relatorio_desligados_html(request):
    data_inicial = request.GET.get("data_inicial")
    data_final = request.GET.get("data_final")

    familias_desligadas = []
    alunos_desligados = []
    erro = None

    if data_inicial and data_final:
        try:
            data_inicial_date = datetime.strptime(data_inicial, "%Y-%m-%d").date()
            data_final_date = datetime.strptime(data_final, "%Y-%m-%d").date()

            inicio_dt = make_aware(datetime.combine(data_inicial_date, time.min))
            fim_dt = make_aware(datetime.combine(data_final_date, time.max))

            family_ct = ContentType.objects.get_for_model(Family)
            aluno_ct = ContentType.objects.get_for_model(Aluno)

            logs_familias = AppLog.objects.filter(
                content_type=family_ct,
                criado_em__range=(inicio_dt, fim_dt),
                acao__icontains="familia_inativada",
            ).select_related("criado_por", "content_type").order_by("-criado_em")

            logs_alunos = AppLog.objects.filter(
                content_type=aluno_ct,
                criado_em__range=(inicio_dt, fim_dt),
                acao__icontains="desligado",
            ).select_related("criado_por", "content_type").order_by("-criado_em")

            for log in logs_familias:
                familia = log.objeto
                familias_desligadas.append({
                    "data": log.criado_em,
                    "inscricao": getattr(familia, "registration_number", "") if familia else "",
                    "familia": getattr(familia, "responsible_name", "") if familia else "",
                    "acao": log.acao,
                    "descricao": log.descricao,
                    "usuario": log.criado_por.username if log.criado_por else "",
                })

            for log in logs_alunos:
                aluno = log.objeto
                alunos_desligados.append({
                    "data": log.criado_em,
                    "aluno": getattr(aluno, "name", "") if aluno else "",
                    "inscricao": (
                        aluno.family.registration_number
                        if aluno and getattr(aluno, "family", None)
                        else ""
                    ),
                    "acao": log.acao,
                    "descricao": log.descricao,
                    "usuario": log.criado_por.username if log.criado_por else "",
                })

        except ValueError:
            erro = "Período inválido. Verifique as datas informadas."
    else:
        erro = "Informe a data inicial e a data final."

    context = {
        "data_inicial": data_inicial,
        "data_final": data_final,
        "familias_desligadas": familias_desligadas,
        "alunos_desligados": alunos_desligados,
        "erro": erro,
    }

    return render(request, "AppLSD/relatorio_desligados_html.html", context)


def _limpar_nome_aba(nome):
    if not nome:
        return "Turma"
    nome = str(nome).replace('/', '-').replace('\\', '-').replace('*', '')
    nome = nome.replace('[', '').replace(']', '').replace(':', '').replace('?', '')
    return nome[:31]


SIGLAS_ATIVIDADES = {
    'JUDÔ': 'JUD',
    'JUDO': 'JUD',
    'CORAL': 'COR',
    'BANDA': 'BAN',
    'TEATRO': 'TEA',
    'KUN FO': 'KFU',
    'KUNG FU': 'KFU',
    'CAPOEIRA': 'CAP',
    'INFORMÁTICA': 'INFO',
    'INFORMATICA': 'INFO',
    'FLAUTA': 'FLA',
    'INGLÊS': 'ING',
    'INGLES': 'ING',
    'HANDEBOL': 'HAN',
    'FUTSAL': 'FUT',
    'VOLEIBOL': 'VOL',
    'VOLEI': 'VOL',
    'EDUCAÇÃO FÍSICA': 'EDF',
    'EDUCACAO FISICA': 'EDF',
}


def normalizar_nome_atividade(nome):
    return (nome or '').strip().upper()


def sigla_atividade(nome):
    nome_normalizado = normalizar_nome_atividade(nome)

    for chave, sigla in SIGLAS_ATIVIDADES.items():
        if chave in nome_normalizado:
            return sigla

    nome_limpo = nome_normalizado.replace(' ', '')
    return nome_limpo[:6] if nome_limpo else 'ATV'


def limpar_nome_aba(nome):
    if not nome:
        return "Turma"
    invalidados = ['\\', '/', '*', '[', ']', ':', '?']
    for ch in invalidados:
        nome = nome.replace(ch, '-')
    return nome[:31]


def obter_logo_path():
    candidatos = [
        'img/logo_integra_lar.png',
        'img/logo_integralar.png',
        'images/logo_integra_lar.png',
        'AppLSD/img/logo_integra_lar.png',
        'AppLSD/img/logo_integralar.png',
    ]

    for caminho in candidatos:
        arquivo = finders.find(caminho)
        if arquivo:
            return arquivo

    return None


def obter_atividades_consolidadas_da_turma(turma):
    atividades_turno = (
        Activity.objects
        .filter(turno=turma.turno)
        .order_by('atividade')
    )

    mapa_siglas = OrderedDict()

    for atividade in atividades_turno:
        nome_original = (atividade.atividade or '').strip()
        sigla = sigla_atividade(nome_original)

        if not sigla:
            continue

        if sigla not in mapa_siglas:
            mapa_siglas[sigla] = {
                'sigla': sigla,
                'nomes_originais': [],
                'ids': set(),
            }

        mapa_siglas[sigla]['ids'].add(atividade.id)

        if nome_original and nome_original not in mapa_siglas[sigla]['nomes_originais']:
            mapa_siglas[sigla]['nomes_originais'].append(nome_original)

    return mapa_siglas


def obter_alunos_da_turma(turma):
    return (
        Aluno.objects
        .filter(turma=turma, status_lsd='Frequentando')
        .select_related('family', 'turma')
        .prefetch_related('atividades')
        .order_by('name')
    )


@login_required
def exportar_alunos_sem_atividades_excel(request):
    turma_id = request.GET.get('turma_id')

    turmas = Turma.objects.select_related('educadora').order_by('turno', 'grupo')
    if turma_id:
        turmas = turmas.filter(id=turma_id)

    wb = Workbook()
    ws_padrao = wb.active
    wb.remove(ws_padrao)

    logo_path = obter_logo_path()

    fill_header = PatternFill(fill_type='solid', start_color='D9EAD3', end_color='D9EAD3')
    fill_info = PatternFill(fill_type='solid', start_color='F4F6F8', end_color='F4F6F8')
    font_header = Font(bold=True, size=11)
    font_titulo = Font(bold=True, size=14)
    font_subtitulo = Font(bold=True, size=11)
    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    thin = Side(border_style='thin', color='A0A0A0')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for idx, turma in enumerate(turmas, start=1):
        nome_aba = limpar_nome_aba(f"{turma.grupo or 'Turma'} - {turma.turno or ''}".strip())
        ws = wb.create_sheet(title=nome_aba or f"Turma {idx}")

        atividades_map = obter_atividades_consolidadas_da_turma(turma)
        siglas = list(atividades_map.keys())

        nome_educadora = ''
        if turma.educadora:
            nome_educadora = turma.educadora.get_full_name() or turma.educadora.first_name or str(turma.educadora)

        total_colunas = 3 + len(siglas) + 2
        ultima_coluna = get_column_letter(total_colunas)

        if logo_path:
            try:
                img = XLImage(logo_path)
                img.width = 85
                img.height = 85
                ws.add_image(img, 'A1')
            except Exception:
                pass

        ws.merge_cells(f'B1:{ultima_coluna}1')
        ws['B1'] = 'LAR SÃO DOMINGOS - RELATÓRIO DE ATIVIDADES'
        ws['B1'].font = font_titulo
        ws['B1'].alignment = align_center

        ws.merge_cells(f'B2:{ultima_coluna}2')
        ws['B2'] = f'Turma: {turma.grupo or "-"}'
        ws['B2'].font = font_subtitulo
        ws['B2'].alignment = align_left
        ws['B2'].fill = fill_info

        ws.merge_cells(f'B3:{ultima_coluna}3')
        ws['B3'] = f'Educadora: {nome_educadora or "-"}'
        ws['B3'].font = font_subtitulo
        ws['B3'].alignment = align_left
        ws['B3'].fill = fill_info

        ws.merge_cells(f'B4:{ultima_coluna}4')
        ws['B4'] = f'Turno: {turma.turno or "-"}   |   Ano Letivo: {turma.ano_letivo}'
        ws['B4'].alignment = align_left
        ws['B4'].fill = fill_info

        if siglas:
            descricao_siglas = ' | '.join(
                f"{sigla} - {', '.join(dados['nomes_originais'])}"
                for sigla, dados in atividades_map.items()
            )
        else:
            descricao_siglas = 'Nenhuma atividade cadastrada para o turno desta turma.'

        ws.merge_cells(f'A6:{ultima_coluna}6')
        ws['A6'] = descricao_siglas
        ws['A6'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        ws['A6'].font = Font(bold=True)

        linha_header = 8
        headers = ['INSC', 'NOME', 'NASC.'] + siglas + ['TOTAL', 'SEM ATIVIDADES']

        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=linha_header, column=col_idx, value=header)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = border

        alunos = obter_alunos_da_turma(turma)
        linha_atual = linha_header + 1
        total_sem_atividades = 0

        for aluno in alunos:
            inscricao = ''
            if aluno.family and getattr(aluno.family, 'registration_number', None):
                inscricao = aluno.family.registration_number
            else:
                inscricao = aluno.id

            siglas_aluno = set()

            for atividade in aluno.atividades.all():
                if atividade.turno != turma.turno:
                    continue

                sigla = sigla_atividade(atividade.atividade)
                if sigla in atividades_map:
                    siglas_aluno.add(sigla)

            total_unico = len(siglas_aluno)
            if total_unico == 0:
                total_sem_atividades += 1

            linha = [
                inscricao,
                aluno.name,
                aluno.birth_date.strftime('%d/%m/%Y') if aluno.birth_date else '',
            ]

            for sigla in siglas:
                linha.append('X' if sigla in siglas_aluno else '')

            linha.append(total_unico)
            linha.append('SIM' if total_unico == 0 else '')

            for col_idx, valor in enumerate(linha, start=1):
                cell = ws.cell(row=linha_atual, column=col_idx, value=valor)
                cell.border = border
                cell.alignment = align_center if col_idx != 2 else align_left

            linha_atual += 1

        ws.cell(row=linha_atual + 1, column=1, value='TOTAL DE ALUNOS').font = font_header
        ws.cell(row=linha_atual + 1, column=2, value=alunos.count())

        ws.cell(row=linha_atual + 2, column=1, value='ALUNOS SEM ATIVIDADES').font = font_header
        ws.cell(row=linha_atual + 2, column=2, value=total_sem_atividades)

        ws.column_dimensions['A'].width = 14
        ws.column_dimensions['B'].width = 40
        ws.column_dimensions['C'].width = 14

        for i in range(4, 4 + len(siglas)):
            ws.column_dimensions[get_column_letter(i)].width = 10

        ws.column_dimensions[get_column_letter(4 + len(siglas))].width = 10
        ws.column_dimensions[get_column_letter(5 + len(siglas))].width = 18

        ws.row_dimensions[1].height = 28
        ws.row_dimensions[6].height = 40
        ws.freeze_panes = f'A{linha_header + 1}'

    if not wb.sheetnames:
        ws = wb.create_sheet(title='Relatório')
        ws['A1'] = 'Nenhuma turma encontrada.'

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="relatorio_alunos_sem_atividades.xlsx"'
    wb.save(response)
    return response


@login_required
def relatorio_alunos_sem_atividades_html(request):
    turma_id = request.GET.get('turma_id') or ''

    turmas = Turma.objects.select_related('educadora').order_by('turno', 'grupo')
    if turma_id:
        turmas = turmas.filter(id=turma_id)

    relatorios = []

    for turma in turmas:
        atividades_map = obter_atividades_consolidadas_da_turma(turma)
        siglas = list(atividades_map.keys())
        alunos = obter_alunos_da_turma(turma)

        nome_educadora = ''
        if turma.educadora:
            nome_educadora = turma.educadora.get_full_name() or turma.educadora.first_name or str(turma.educadora)

        linhas = []
        total_sem_atividades = 0

        for aluno in alunos:
            inscricao = ''
            if aluno.family and getattr(aluno.family, 'registration_number', None):
                inscricao = aluno.family.registration_number
            else:
                inscricao = aluno.id

            siglas_aluno = set()

            for atividade in aluno.atividades.all():
                if atividade.turno != turma.turno:
                    continue

                sigla = sigla_atividade(atividade.atividade)
                if sigla in atividades_map:
                    siglas_aluno.add(sigla)

            total_unico = len(siglas_aluno)
            sem_atividades = total_unico == 0

            if sem_atividades:
                total_sem_atividades += 1

            marcacoes = []
            for sigla in siglas:
                marcacoes.append({
                    'sigla': sigla,
                    'marcado': sigla in siglas_aluno,
                })

            linhas.append({
                'inscricao': inscricao,
                'nome': aluno.name,
                'nascimento': aluno.birth_date.strftime('%d/%m/%Y') if aluno.birth_date else '',
                'marcacoes': marcacoes,
                'total': total_unico,
                'sem_atividades': sem_atividades,
            })

        atividades_legenda = []
        for sigla, dados in atividades_map.items():
            atividades_legenda.append({
                'sigla': sigla,
                'descricao': ', '.join(dados['nomes_originais']),
            })

        relatorios.append({
            'turma': turma,
            'nome_educadora': nome_educadora,
            'atividades': atividades_legenda,
            'siglas': siglas,
            'linhas': linhas,
            'total_alunos': alunos.count(),
            'total_sem_atividades': total_sem_atividades,
        })

    context = {
        'relatorios': relatorios,
        'turma_id': turma_id,
    }
    return render(request, 'AppLSD/relatorio_alunos_sem_atividades.html', context)